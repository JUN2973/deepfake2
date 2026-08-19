from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "docs/회원가입_로그인_코드설명_프린트용.docx"


def set_run_font(run, name="맑은 고딕", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, name="맑은 고딕", size=None, color=None, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=9.5, bold=bold)
    if fill:
        shade_cell(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True, fill="E8EEF5")
        set_cell_width(hdr[i], widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            set_cell_width(cells[i], widths[i])
    doc.add_paragraph()
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
    p.paragraph_format.first_line_indent = Inches(-0.1)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def add_step(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title + " ")
    set_run_font(run, size=10.5, bold=True, color="1F4D78")
    run = p.add_run(body)
    set_run_font(run, size=10.5)


def add_code_path(doc, label, path):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(label + ": ")
    set_run_font(run, size=9.5, bold=True)
    run = p.add_run(path)
    set_run_font(run, name="Consolas", size=9)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    set_run_font(run, size=10.5, bold=True, color="1F3A5F")
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(body)
    set_run_font(run, size=10)
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    styles = doc.styles
    set_style_font(styles["Normal"], size=10.5)
    styles["Normal"].paragraph_format.line_spacing = 1.18
    styles["Normal"].paragraph_format.space_after = Pt(5)
    set_style_font(styles["Heading 1"], size=16, color="2E74B5", bold=True)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    set_style_font(styles["Heading 2"], size=13, color="2E74B5", bold=True)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    set_style_font(styles["Heading 3"], size=12, color="1F4D78", bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("DeepScan 회원가입/로그인 코드 설명")
    set_run_font(run, size=8.5, color="666666")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("회원가입 및 로그인 코드 설명서")
    set_run_font(run, size=20, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("DeepScan 프로젝트 | 프린트 및 발표 정리용")
    set_run_font(run, size=10, color="555555")

    add_callout(
        doc,
        "핵심 한 줄 요약",
        "이 프로젝트는 JSP 화면에서 fetch로 인증 API를 호출하고, AuthApiController와 AuthService가 MyBatis Mapper를 통해 MariaDB의 app_user, email_auth 테이블을 조회/저장하는 서버 세션 기반 인증 구조입니다.",
    )

    doc.add_heading("1. 전체 구조", level=1)
    add_bullet(doc, "화면 진입은 AuthPageController가 담당합니다. /signup, /login 요청을 JSP 화면으로 연결합니다.")
    add_bullet(doc, "회원가입/로그인 실제 처리는 /api/v1/auth 하위 API를 제공하는 AuthApiController가 담당합니다.")
    add_bullet(doc, "비즈니스 검증, 비밀번호 해시, 이메일 인증 확인은 AuthService와 EmailAuthService에서 처리합니다.")
    add_bullet(doc, "DB 접근은 IUserMapper, IEmailAuthMapper와 MyBatis XML SQL에서 처리합니다.")
    add_bullet(doc, "로그인 상태는 JWT가 아니라 HttpSession의 USER_ID, USER_NAME, USER_EMAIL로 관리합니다.")

    doc.add_heading("2. 주요 코드 경로", level=1)
    add_table(
        doc,
        ["구분", "파일 경로", "역할"],
        [
            ["화면 이동", "src/main/java/kopo/poly/controller/AuthPageController.java", "/signup, /login, /find-account, /mypage 화면 반환"],
            ["로그인 화면", "src/main/webapp/WEB-INF/views/login.jsp", "로그인 폼, 일반 로그인 API 호출, 구글 로그인 진입"],
            ["회원가입 화면", "src/main/webapp/WEB-INF/views/signup.jsp", "중복확인, 인증번호 발송/확인, 최종 가입 API 호출"],
            ["인증 API", "src/main/java/kopo/poly/controller/api/AuthApiController.java", "로그인, 회원가입, 로그아웃, 계정 관련 REST API 처리"],
            ["인증 서비스", "src/main/java/kopo/poly/service/impl/AuthService.java", "입력값 검증, 회원 조회, 비밀번호 비교/저장, 세션용 사용자 정보 반환"],
            ["이메일 인증", "src/main/java/kopo/poly/service/impl/EmailAuthService.java", "6자리 인증번호 생성, 저장, 메일 발송, 검증"],
            ["회원 SQL", "src/main/resources/mapper/UserMapper.xml", "app_user 테이블 조회/저장/수정/삭제 SQL"],
            ["인증번호 SQL", "src/main/resources/mapper/EmailAuthMapper.xml", "email_auth 테이블 인증번호 저장/조회/검증 완료 SQL"],
        ],
        [Inches(1.1), Inches(3.05), Inches(2.35)],
    )

    doc.add_heading("3. 회원가입 흐름", level=1)
    add_step(doc, "1단계: 이메일 중복 확인", "signup.jsp에서 /api/v1/auth/signup/check-email을 호출합니다. AuthService.checkEmail()은 이메일 형식을 정리하고 UserMapper.existsByEmail()로 app_user에 같은 이메일이 있는지 확인합니다.")
    add_code_path(doc, "화면", "src/main/webapp/WEB-INF/views/signup.jsp:448")
    add_code_path(doc, "컨트롤러", "src/main/java/kopo/poly/controller/api/AuthApiController.java:143")
    add_code_path(doc, "SQL", "src/main/resources/mapper/UserMapper.xml:13")

    add_step(doc, "2단계: 인증번호 발송", "중복이 없으면 /api/v1/auth/signup/send-code를 호출합니다. EmailAuthService가 SecureRandom으로 6자리 인증번호를 만들고, 만료 시간을 5분으로 설정한 뒤 email_auth 테이블에 저장하고 메일을 보냅니다.")
    add_code_path(doc, "화면", "src/main/webapp/WEB-INF/views/signup.jsp:494")
    add_code_path(doc, "서비스", "src/main/java/kopo/poly/service/impl/EmailAuthService.java:52")
    add_code_path(doc, "SQL", "src/main/resources/mapper/EmailAuthMapper.xml:12")

    add_step(doc, "3단계: 인증번호 확인", "사용자가 입력한 인증번호를 /api/v1/auth/signup/verify-code로 보냅니다. 서버는 가장 최근 인증번호를 조회하고, 만료 여부와 코드 일치 여부를 검사한 뒤 verified 값을 Y로 변경합니다.")
    add_code_path(doc, "화면", "src/main/webapp/WEB-INF/views/signup.jsp:529")
    add_code_path(doc, "컨트롤러", "src/main/java/kopo/poly/controller/api/AuthApiController.java:176")
    add_code_path(doc, "SQL", "src/main/resources/mapper/EmailAuthMapper.xml:32, 55")

    add_step(doc, "4단계: 최종 회원가입", "이메일 인증이 완료된 상태에서 /api/v1/auth/signup/complete를 호출합니다. AuthService.completeSignup()은 필수값과 인증 여부를 다시 확인하고, 비밀번호를 BCrypt로 해시 처리해 app_user 테이블에 저장합니다.")
    add_code_path(doc, "화면", "src/main/webapp/WEB-INF/views/signup.jsp:597")
    add_code_path(doc, "서비스", "src/main/java/kopo/poly/service/impl/AuthService.java:246")
    add_code_path(doc, "저장 SQL", "src/main/resources/mapper/UserMapper.xml:18")

    doc.add_heading("4. 로그인 흐름", level=1)
    add_step(doc, "1단계: 로그인 API 호출", "login.jsp에서 사용자가 입력한 email/password를 JSON으로 만들어 /api/v1/auth/login에 POST 요청합니다.")
    add_code_path(doc, "화면", "src/main/webapp/WEB-INF/views/login.jsp:169")
    add_code_path(doc, "요청 DTO", "src/main/java/kopo/poly/dto/LoginRequestDTO.java")

    add_step(doc, "2단계: 회원 조회", "AuthApiController.login()이 요청을 받고 AuthService.login()으로 넘깁니다. 서비스는 userMapper.selectUserByEmail(email)로 app_user에서 회원을 찾습니다.")
    add_code_path(doc, "컨트롤러", "src/main/java/kopo/poly/controller/api/AuthApiController.java:52")
    add_code_path(doc, "서비스", "src/main/java/kopo/poly/service/impl/AuthService.java:55")
    add_code_path(doc, "조회 SQL", "src/main/resources/mapper/UserMapper.xml:30")

    add_step(doc, "3단계: 비밀번호 비교", "DB에는 원문 비밀번호가 아니라 password_hash가 저장됩니다. BCrypt 해시는 PASSWORD_ENCODER.matches()로 비교하고, 예전 SHA-256 해시도 로그인 시 호환합니다.")
    add_code_path(doc, "BCrypt 설정", "src/main/java/kopo/poly/service/impl/AuthService.java:39")
    add_code_path(doc, "비밀번호 비교", "src/main/java/kopo/poly/service/impl/AuthService.java:384")
    add_code_path(doc, "SHA-256 호환", "src/main/java/kopo/poly/service/impl/AuthService.java:395")

    add_step(doc, "4단계: 세션 저장", "로그인 성공 시 request.changeSessionId()로 세션 ID를 교체하고, USER_ID, USER_NAME, USER_EMAIL을 세션에 저장합니다.")
    add_code_path(doc, "세션 저장", "src/main/java/kopo/poly/controller/api/AuthApiController.java:58")
    add_code_path(doc, "세션 조회 유틸", "src/main/java/kopo/poly/util/SessionUtil.java:20")

    doc.add_heading("5. 관련 테이블", level=1)
    add_table(
        doc,
        ["테이블", "주요 컬럼", "역할"],
        [
            ["app_user", "id, name, email, password_hash, phone_number, address, oauth_provider", "회원 기본 정보와 비밀번호 해시, 소셜 로그인 연결 정보를 저장합니다."],
            ["email_auth", "id, email, auth_code, verified, expires_at, created_at, updated_at", "회원가입/비밀번호 재설정용 이메일 인증번호와 검증 상태를 저장합니다."],
        ],
        [Inches(1.25), Inches(2.6), Inches(2.65)],
    )

    doc.add_heading("6. 보안 포인트", level=1)
    add_bullet(doc, "비밀번호는 원문 저장 없이 BCrypt 해시로 저장합니다.")
    add_bullet(doc, "예전 SHA-256 비밀번호도 로그인 성공 시 BCrypt로 자동 전환합니다.")
    add_bullet(doc, "로그인 성공 시 changeSessionId()를 호출해 세션 고정 공격 위험을 줄입니다.")
    add_bullet(doc, "회원가입 최종 단계에서 이메일 인증 여부와 이메일 중복 여부를 서버에서 다시 확인합니다.")
    add_bullet(doc, "이메일 인증번호는 5분 만료이며, 새 인증번호를 발급하면 이전 인증번호를 무효화합니다.")

    doc.add_heading("7. 구글 로그인 흐름", level=1)
    add_bullet(doc, "login.jsp의 구글 로그인 버튼은 /oauth2/authorization/google로 이동합니다.")
    add_bullet(doc, "GoogleOAuthController가 state 값을 세션에 저장한 뒤 Google 인증 화면으로 redirect합니다.")
    add_bullet(doc, "callback에서는 state를 검증하고, authorization code를 access token으로 교환합니다.")
    add_bullet(doc, "GoogleOAuthService가 userinfo에서 sub, email, name을 읽어 기존 계정 연결 또는 소셜 회원 생성을 수행합니다.")
    add_bullet(doc, "성공 후 일반 로그인과 동일하게 USER_ID, USER_NAME, USER_EMAIL을 세션에 저장합니다.")

    doc.add_heading("8. 발표용 요약 문장", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(
        "회원가입은 이메일 중복 확인, 인증번호 발송, 인증번호 검증, 최종 가입 4단계로 구성했습니다. "
        "화면에서는 signup.jsp가 fetch로 API를 호출하고, 서버에서는 AuthApiController가 요청을 받아 AuthService로 전달합니다. "
        "인증번호는 EmailAuthService가 생성해 email_auth 테이블에 저장하고 메일로 발송합니다. "
        "인증이 완료된 사용자만 app_user 테이블에 저장되며, 비밀번호는 BCrypt로 해시 처리합니다."
    )
    set_run_font(run, size=10.5)
    p = doc.add_paragraph()
    run = p.add_run(
        "로그인은 login.jsp에서 /api/v1/auth/login으로 이메일과 비밀번호를 보내고, AuthService.login()이 app_user에서 회원을 조회한 뒤 비밀번호 해시를 비교합니다. "
        "성공하면 AuthApiController가 세션 ID를 교체하고 USER_ID, USER_NAME, USER_EMAIL을 세션에 저장합니다. "
        "따라서 이 프로젝트는 JWT가 아니라 서버 세션 기반 인증 방식입니다."
    )
    set_run_font(run, size=10.5)

    doc.save(OUT)


if __name__ == "__main__":
    build()
