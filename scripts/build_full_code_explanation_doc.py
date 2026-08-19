from pathlib import Path
import re

from docx import Document
from docx.shared import Pt, RGBColor


ROOT = Path(r"C:\SpringBootWorks\deepfake2")
OUTPUT = ROOT / "docs" / "딥페이크_프로젝트_전체코드_경로_메인기능_설명서.docx"


INCLUDE_PATTERNS = {
    ".java", ".xml", ".jsp", ".properties", ".yml", ".yaml", ".py"
}


MANUAL = {
    "Deepfake2Application.java": "Spring Boot 시작 클래스입니다. main()에서 애플리케이션을 실행하고 @MapperScan으로 kopo.poly.mapper 패키지의 MyBatis Mapper 인터페이스를 스캔합니다.",
    "AppConfig.java": "RestTemplate 공통 Bean을 등록하고 HTTP connect/read timeout을 설정합니다. 외부 API 호출 코드들이 이 Bean을 사용합니다.",
    "WebConfig.java": "정적 리소스와 업로드 파일 접근 경로를 설정합니다. 저장된 이미지가 화면에서 publicUrl로 보이게 하는 데 연결됩니다.",
    "AppUserSchemaInitializer.java": "app_user 테이블이 없을 때 생성/보정하는 초기화 코드입니다. 회원, 소셜 로그인, 프로필 정보를 저장하는 기반입니다.",
    "EmailAuthSchemaInitializer.java": "email_auth 테이블을 준비합니다. 회원가입과 비밀번호 재설정 인증번호 저장에 사용됩니다.",
    "VerificationSchemaInitializer.java": "verification 테이블을 준비합니다. 이미지 업로드 분석 결과와 원본 API 응답을 저장합니다.",
    "DetectionSchemaInitializer.java": "verification_record 테이블을 준비합니다. Reality Defender 판별 기록 화면에서 사용하는 레거시/별도 기록 저장 구조입니다.",
    "SessionUtil.java": "세션에서 USER_ID를 읽고, 검증 기록 접근 권한을 확인하고, 최근 검증 ID를 기억하는 공통 유틸입니다.",
    "HashUtil.java": "비밀번호를 SHA-256으로 해시하는 유틸입니다. 회원가입, 로그인, 비밀번호 변경에서 사용됩니다.",
    "CommonUtil.java": "문자열 null 처리 같은 공통 보조 함수를 제공하는 유틸입니다.",
    "ApiResponse.java": "REST API 응답 형식을 success/data/error 구조로 통일하는 공통 DTO입니다. ok(), fail()로 성공/실패 응답을 만듭니다.",
    "VerifyService.java": "메인 기능의 핵심 서비스입니다. 이미지 검증, 파일 저장, 사전 이미지 품질 검사, 딥페이크 분석 API 호출, 히트맵 병합, DB 저장, 상세 분석 JSON 생성까지 담당합니다.",
    "RealityDefenderClient.java": "Reality Defender 실제 API를 호출하는 딥페이크 분석 클라이언트입니다. 파일 업로드/분석 요청/결과 변환을 담당합니다.",
    "ImdDeepfakeClient.java": "로컬 Python IMD 서비스에 이미지를 multipart로 보내 딥페이크 분석 결과를 받는 클라이언트입니다.",
    "ImdHeatmapClient.java": "IMD 서비스에서 히트맵, 오버레이, 의심 영역 정보를 받아 VerifyService의 결과에 보강합니다.",
    "ImdLocalServiceManager.java": "IMD 로컬 Python 서버 자동 실행과 상태 확인을 담당합니다. 설정값에 따라 로컬 분석 서비스를 띄울 수 있습니다.",
    "DummyDeepfakeClient.java": "실제 API 없이 테스트용 판정 결과를 반환하는 더미 딥페이크 클라이언트입니다.",
    "IObjectStorageService.java": "파일 저장소 인터페이스입니다. 업로드 결과로 objectKey와 publicUrl을 반환하도록 약속합니다.",
    "DummyObjectStorageService.java": "로컬 uploads 폴더에 파일을 저장하는 구현체입니다. 개발 환경 기본 저장소 역할입니다.",
    "S3ObjectStorageService.java": "Amazon S3에 파일을 업로드하는 구현체입니다. 설정값으로 storage type을 바꿨을 때 사용할 수 있습니다.",
    "VerificationApiController.java": "검증 기록 REST API입니다. 이미지 업로드 생성, 상세 조회, 목록 조회, 삭제를 /api/v1/verifications 경로로 처리합니다.",
    "UploadController.java": "JSP 업로드 화면에서 /upload로 들어온 파일을 받아 VerifyService로 넘기고 분석 후 상세 결과 화면으로 이동시킵니다.",
    "VerifyController.java": "레거시 /verify 업로드 요청을 처리합니다. 업로드 파일을 분석 흐름으로 연결합니다.",
    "ResultController.java": "검증 결과 상세 화면 /detail/{id}를 반환합니다. 저장된 분석 결과를 화면 모델로 넘깁니다.",
    "DetailController.java": "레거시 상세 조회 /legacy/detail/{id} 화면을 처리합니다.",
    "HistoryController.java": "로그인 사용자의 검증 기록 목록 화면 /history를 처리합니다.",
    "AnalyzingController.java": "분석 중 화면 /analyzing을 반환합니다.",
    "RealityDefenderController.java": "Reality Defender 전용 판별 화면 /detect, 결과, 히스토리 흐름을 담당합니다.",
    "RealityDefenderService.java": "Reality Defender 전용 업로드, 판별, 기록 저장, 결과 파싱을 처리하는 서비스입니다.",
    "IRealityDefenderService.java": "Reality Defender 기능이 제공해야 하는 메서드 계약입니다.",
    "IRealityDefenderMapper.java": "verification_record 관련 SQL 호출 인터페이스입니다.",
    "RealityDefenderMapper.xml": "Reality Defender 판별 기록 저장/조회 SQL을 담고 있습니다.",
    "AuthApiController.java": "회원 API 핵심 컨트롤러입니다. 로그인, 회원가입, 아이디 찾기, 비밀번호 재설정, 프로필 수정, 비밀번호 변경, 로그아웃, 회원 탈퇴를 처리합니다.",
    "AuthPageController.java": "회원 관련 JSP 화면으로 이동합니다. /signup, /login, /find-account, /mypage를 반환합니다.",
    "EmailAuthController.java": "이메일 인증번호 발송/검증 API입니다. /auth/email/send, /auth/email/verify를 처리합니다.",
    "EmailAuthService.java": "인증번호 생성, 저장, 메일 발송, 인증번호 검증, 인증 완료 여부 확인을 담당합니다.",
    "MailService.java": "JavaMailSender로 실제 메일을 발송합니다.",
    "GoogleOAuthController.java": "구글 OAuth 시작 요청과 callback 요청을 처리하고, 로그인 성공 시 USER_ID/USER_NAME/USER_EMAIL 세션을 저장합니다.",
    "GoogleOAuthService.java": "구글 토큰 교환, 사용자 정보 조회, 기존 계정 연결 또는 소셜 회원 생성을 담당합니다.",
    "IUserMapper.java": "회원 테이블 app_user 조회/저장/수정/삭제 SQL을 호출하는 MyBatis 인터페이스입니다.",
    "UserMapper.xml": "로그인, 회원가입, 아이디 찾기, 비밀번호 변경, 프로필 수정, 회원 탈퇴 관련 SQL을 담고 있습니다.",
    "IEmailAuthMapper.java": "email_auth 테이블 조회/저장/수정 SQL 인터페이스입니다.",
    "EmailAuthMapper.xml": "이메일 인증번호 저장, 인증 확인, 인증 완료 처리 SQL을 담고 있습니다.",
    "NewsController.java": "뉴스 화면 컨트롤러입니다. /news 목록, /news/refresh 새로고침, /news/{id} 상세 화면을 처리합니다.",
    "NewsService.java": "Naver News API 호출, 기사 정제, 딥페이크 관련 필터링, 중복 제거, MongoDB/메모리 캐시 저장, 상세 기사 보강을 담당합니다.",
    "NewsCacheRepository.java": "MongoDB news_cache 컬렉션 접근 Repository입니다.",
    "NewsCacheDocument.java": "MongoDB에 저장되는 뉴스 캐시 문서 구조입니다.",
    "NaverNewsResponseDTO.java": "Naver News API 전체 응답 DTO입니다.",
    "NaverNewsItemDTO.java": "Naver News API 기사 1건 DTO입니다.",
    "NewsViewDTO.java": "뉴스 상세 화면에 보여줄 제목, 요약, 언론사, 날짜, 링크, 태그를 담는 화면용 DTO입니다.",
    "CommunityController.java": "커뮤니티 JSP 화면 이동 컨트롤러입니다. 목록, 작성, 상세 화면으로 이동합니다.",
    "CommunityApiController.java": "커뮤니티 REST API입니다. 게시글 목록/상세/작성/수정/삭제, 좋아요, 댓글 목록/작성/수정/삭제, 댓글 좋아요를 처리합니다.",
    "CommunityService.java": "게시글과 댓글의 비즈니스 로직입니다. 작성자 권한 확인, 조회수 증가, 좋아요 토글, 댓글 트리 구성을 담당합니다.",
    "ICommunityMapper.java": "커뮤니티 게시글/댓글/좋아요 SQL 인터페이스입니다.",
    "CommunityMapper.xml": "게시글, 댓글, 좋아요의 CRUD와 조회수 증가 SQL을 담고 있습니다.",
    "FaqController.java": "FAQ 화면 /faq를 반환합니다.",
    "ReportController.java": "신고 안내 화면 /report를 반환합니다.",
    "HomeController.java": "메인 화면 / 또는 /home을 반환합니다.",
    "DetailService.java": "상세 결과 조회 기능을 위한 서비스입니다.",
    "HistoryService.java": "검증 기록 목록 조회를 위한 서비스입니다.",
    "IDetailMapper.java": "상세 조회 SQL 인터페이스입니다.",
    "DetailMapper.xml": "상세 조회 SQL을 담고 있습니다.",
    "IHistoryMapper.java": "검증 이력 조회 SQL 인터페이스입니다.",
    "HistoryMapper.xml": "검증 기록 목록 조회 SQL을 담고 있습니다.",
    "IVerifyMapper.java": "verification 테이블 저장/조회/삭제 SQL 인터페이스입니다.",
    "VerifyMapper.xml": "검증 결과 insert, 상세 select, 목록 select, 삭제 SQL을 담고 있습니다.",
    "application.properties": "프로젝트 전체 설정 파일입니다. JSP 경로, 포트, 업로드 제한, 저장소, DB, 메일, Reality Defender, IMD, Naver News, MongoDB, Google OAuth2 설정을 관리합니다.",
    "app.py": "Python 기반 IMD 로컬 분석 서비스입니다. Spring Boot의 ImdDeepfakeClient/ImdHeatmapClient가 HTTP로 호출합니다.",
    "docker-compose.imd.prod.yml": "운영 또는 배포 환경에서 IMD 서비스를 컨테이너로 실행하기 위한 compose 설정입니다.",
    "docker-compose.imd.yml": "개발 환경에서 IMD 관련 서비스를 실행하기 위한 compose 설정입니다.",
}


JSP_MANUAL = {
    "deepfake-upload.jsp": "딥페이크 이미지 업로드 화면입니다. 파일 선택, 미리보기, 업로드 요청을 통해 /upload 또는 검증 API 흐름으로 들어갑니다.",
    "deepfake-analyzing.jsp": "분석 진행 중임을 보여주는 화면입니다. 업로드 후 결과가 준비될 때까지 사용자가 기다리는 중간 화면입니다.",
    "deepfake-result.jsp": "딥페이크 분석 결과 화면입니다. 판정, 점수, 설명, 원본 이미지와 분석 정보를 보여줍니다.",
    "deepfake-history.jsp": "사용자의 검증 기록 목록 화면입니다. 과거 분석 결과를 다시 확인하는入口입니다.",
    "result.jsp": "결과 또는 레거시 결과 표시 화면입니다.",
    "login.jsp": "로그인 화면입니다. 일반 로그인 API와 구글 로그인 진입을 연결합니다.",
    "signup.jsp": "회원가입 화면입니다. 이메일 중복 확인, 인증번호 발송/검증, 최종 가입 API를 호출합니다.",
    "account-find.jsp": "아이디 찾기와 비밀번호 재설정 화면입니다. 이름/전화번호 조회, 계정 확인, 인증번호 발송/검증, 새 비밀번호 저장으로 연결됩니다.",
    "user-profile.jsp": "마이페이지 화면입니다. 내 정보 조회, 프로필 수정, 비밀번호 변경, 로그아웃, 회원 탈퇴 API를 호출합니다.",
    "news-list.jsp": "뉴스 목록 화면입니다. 검색어, 카테고리, 정렬, 페이지 이동, 새로고침을 처리합니다.",
    "news-detail.jsp": "뉴스 상세 화면입니다. 기사 요약과 원문 링크 이동을 제공합니다.",
    "community-list.jsp": "커뮤니티 게시글 목록 화면입니다. 목록 조회, 검색, 상세/작성 화면 이동을 처리합니다.",
    "community-post-create.jsp": "게시글 작성/수정 화면입니다. 게시글 저장 API로 연결됩니다.",
    "community-post-detail.jsp": "게시글 상세 화면입니다. 댓글, 대댓글, 좋아요, 수정/삭제 API와 연결됩니다.",
    "faq.jsp": "자주 묻는 질문 화면입니다. 질문을 펼쳐 답변을 보는 정적/화면 기능입니다.",
    "report.jsp": "신고 안내 화면입니다. 딥페이크 피해 신고 방법을 안내합니다.",
}


DTO_MANUAL = {
    "LoginRequestDTO.java": "로그인 요청에서 email/password를 받습니다.",
    "SignupCompleteRequestDTO.java": "최종 회원가입 단계의 email, password, name, phone, address 값을 받습니다.",
    "SendCodeRequestDTO.java": "이메일 인증번호 발송 요청에서 email을 받습니다.",
    "VerifyCodeRequestDTO.java": "비밀번호 재설정/회원가입 인증번호 검증 요청에서 email과 code를 받습니다.",
    "EmailAuthSendRequestDTO.java": "이메일 인증 API의 발송 요청 DTO입니다.",
    "EmailAuthVerifyRequestDTO.java": "이메일 인증 API의 검증 요청 DTO입니다.",
    "EmailAuthResponseDTO.java": "이메일 인증 성공 여부와 메시지를 반환합니다.",
    "FindIdRequestDTO.java": "아이디 찾기 요청에서 이름과 전화번호를 받습니다.",
    "FindPasswordRequestDTO.java": "비밀번호 재설정 인증번호 발송 요청에서 이메일과 이름을 받습니다.",
    "ResetPasswordRequestDTO.java": "새 비밀번호 저장 요청에서 email/password를 받습니다.",
    "VerifyDTO.java": "verification 테이블의 검증 결과 데이터입니다. 파일명, URL, 판정, 점수, API 원본 응답, userId 등을 담습니다.",
    "VerificationCreateResponseDTO.java": "검증 생성 API가 반환하는 id, 파일명, 저장 경로, publicUrl, 판정, 점수, 생성일 DTO입니다.",
    "VerificationRecordDTO.java": "Reality Defender 별도 검증 기록 DTO입니다.",
    "DetectionResultDTO.java": "판별 결과 화면에 필요한 판정/점수/메시지 계열 데이터를 담습니다.",
    "DeepfakeResultDTO.java": "딥페이크 분석 클라이언트가 반환하는 verdict, score, raw JSON DTO입니다.",
    "DetailDTO.java": "상세 조회 화면에서 사용할 검증 상세 DTO입니다.",
    "HistoryDTO.java": "검증 기록 목록 화면에서 사용할 이력 DTO입니다.",
    "SuspiciousRegionDTO.java": "의심 영역 좌표나 confidence 정보를 담는 DTO입니다.",
    "MailDTO.java": "메일 발송에 필요한 수신자, 제목, 내용 정보를 담습니다.",
    "CommunityCreateRequestDTO.java": "게시글 작성/수정 요청에서 topic, title, content를 받습니다.",
    "CommunityPostDTO.java": "커뮤니티 게시글 상세/목록 데이터입니다.",
    "CommunityPostLikeResponseDTO.java": "게시글 좋아요 토글 후 좋아요 여부와 개수를 반환합니다.",
    "CommunityCommentCreateRequestDTO.java": "댓글/대댓글 작성 요청 데이터를 받습니다.",
    "CommunityCommentUpdateRequestDTO.java": "댓글 수정 요청 데이터를 받습니다.",
    "CommunityCommentDTO.java": "댓글과 대댓글 트리 데이터를 담습니다.",
    "CommunityCommentLikeResponseDTO.java": "댓글 좋아요 토글 결과를 반환합니다.",
}


FEATURE_FLOWS = [
    ("메인 딥페이크 분석 기능",
     [
         "진입 화면: 사용자는 / 또는 /home에서 메인 화면을 보고, 업로드 화면 deepfake-upload.jsp로 이동합니다.",
         "업로드 요청: 화면에서 파일을 선택하면 /upload 또는 /api/v1/verifications로 multipart/form-data 요청이 들어갑니다.",
         "Controller: UploadController 또는 VerificationApiController가 request.getFile(\"file\")로 파일을 꺼내 VerifyService.createVerification(file, userId)를 호출합니다.",
         "파일 검증: VerifyService.validateImageOnly()가 MIME 타입, 확장자, 32MB 이하 용량을 검사합니다.",
         "저장 경로 생성: 날짜와 UUID, 원본 파일명을 조합해 objectKey를 만들고 IObjectStorageService.uploadPublic()으로 로컬 uploads 또는 S3에 저장합니다.",
         "사전 이미지 분석: ImageIO로 이미지 크기, 흐림 정도, 피부색 영역, 사람이 너무 많거나 얼굴이 너무 작은지 등을 확인합니다. 분석에 부적절하면 NOT_APPLICABLE 결과를 만듭니다.",
         "외부 분석 호출: 적합한 이미지면 IDeepfakeClient.analyze(temp)를 호출합니다. 설정에 따라 RealityDefenderClient, ImdDeepfakeClient, DummyDeepfakeClient 중 하나가 동작합니다.",
         "히트맵 보강: Reality Defender 결과를 받은 경우 ImdHeatmapClient.generateHeatmap(temp)를 추가 호출해 processedHeatmap, overlay, suspicious regions를 raw JSON에 병합합니다.",
         "DB 저장: VerifyDTO에 파일 정보, verdict, score, apiProvider, apiRaw, userId를 담아 IVerifyMapper.insertVerification()으로 verification 테이블에 저장합니다.",
         "결과 조회: ResultController/DetailController 또는 VerificationApiController.detail()이 VerifyService.getOne(id)을 호출합니다.",
         "상세 분석 생성: VerifyService.enrichVerification()이 api_raw를 파싱해서 분석 출처, 히트맵, 의심 영역, 증상, enhancedExplanation을 만들어 화면에서 쓰기 쉽게 바꿉니다.",
         "화면 출력: deepfake-result.jsp 또는 상세 화면에서 원본 이미지, 판정, 점수, 요약 문구, 근거, 주의사항, 히트맵을 보여줍니다.",
     ]),
    ("회원/로그인 기능",
     [
         "화면 진입: /login은 login.jsp, /signup은 signup.jsp, /find-account는 account-find.jsp, /mypage는 user-profile.jsp로 들어갑니다.",
         "일반 로그인: login.jsp가 /api/v1/auth/login으로 email/password를 보내고 AuthApiController가 IUserMapper.selectUserByEmail()로 회원을 조회합니다.",
         "비밀번호 비교: 입력 비밀번호를 HashUtil.sha256()으로 해시해서 DB 값과 비교하고, 성공하면 세션에 USER_ID, USER_NAME, USER_EMAIL을 저장합니다.",
         "회원가입: check-email, send-code, verify-code, complete 순서로 진행됩니다. 이메일 인증이 완료된 사용자만 app_user에 저장됩니다.",
         "비밀번호 재설정: 이메일과 이름으로 회원을 확인하고 인증번호를 보낸 뒤, 인증 완료 상태일 때 새 비밀번호를 해시해 저장합니다.",
         "구글 로그인: GoogleOAuthController가 OAuth 시작/callback을 처리하고 GoogleOAuthService가 구글 sub/email/name을 읽어 회원 연결 또는 소셜 회원 생성을 수행합니다.",
         "마이페이지: /api/v1/auth/me, /profile, /password, /account, /logout API로 내 정보 조회, 수정, 비밀번호 변경, 탈퇴, 로그아웃을 처리합니다.",
     ]),
    ("뉴스 기능",
     [
         "진입 경로: /news 요청은 NewsController.newsPage()로 들어가 news-list.jsp를 반환합니다.",
         "검색/정렬: query, category, sort, page 파라미터를 읽고 기본 검색어는 딥페이크로 둡니다.",
         "서비스 처리: NewsService.getNews(query, sort)가 로컬 메모리 캐시, MongoDB 캐시, Naver News API 순서로 데이터를 찾습니다.",
         "API 호출: Naver News API에는 X-Naver-Client-Id와 X-Naver-Client-Secret 헤더를 넣어 GET 요청합니다.",
         "정제/필터: 제목과 설명에서 HTML을 제거하고, 딥페이크 관련 키워드 필터링, 영어 기사 제외, 중복 기사 제거를 수행합니다.",
         "캐시 저장: 정제된 뉴스는 NewsCacheDocument로 만들어 MongoDB news_cache에 저장하고, 메모리 캐시에도 올립니다.",
         "상세 보기: /news/{id}는 MongoDB 또는 로컬 캐시에서 기사 정보를 찾고 NewsViewDTO로 바꿔 news-detail.jsp에 전달합니다.",
     ]),
    ("커뮤니티 기능",
     [
         "화면 경로: /community, /community/create, /community/{id}가 JSP 화면으로 들어가는 경로입니다.",
         "API 경로: /api/v1/community/posts가 게시글 목록/작성/상세/수정/삭제의 기준 경로입니다.",
         "게시글 작성: CommunityApiController.create()가 세션 USER_ID/USER_NAME을 확인하고 title/content/topic을 검증한 뒤 CommunityService.createPost()를 호출합니다.",
         "게시글 수정/삭제: Service에서 DB의 작성자 userId와 세션 userId가 같은지 확인한 뒤 Mapper를 호출합니다.",
         "댓글 처리: comments API가 댓글을 조회/작성/수정/삭제하고, CommunityService.buildCommentTree()가 parentId 기준으로 댓글/대댓글 구조를 만듭니다.",
         "좋아요 처리: 같은 API를 다시 누르면 insert/delete가 토글되어 좋아요 추가와 취소를 처리합니다.",
     ]),
    ("검증 기록/결과 조회 기능",
     [
         "목록 진입: /history 화면 또는 GET /api/v1/verifications가 로그인 사용자의 검증 기록을 가져옵니다.",
         "권한 확인: 세션 userId와 기록의 userId를 비교해서 다른 사용자의 기록 접근을 막습니다.",
         "상세 조회: /detail/{id}, /api/v1/verifications/{id}가 VerifyService.getOne(id)을 호출하고 상세 분석 결과를 화면/API로 반환합니다.",
         "삭제: DELETE /api/v1/verifications/{id}는 userId 조건으로 삭제해 본인 기록만 지울 수 있게 합니다.",
     ]),
]


def read_text(path):
    return path.read_text(encoding="utf-8", errors="ignore")


def files():
    roots = [ROOT / "src", ROOT / "imd-service", ROOT / "deploy", ROOT]
    found = []
    excluded_parts = {
        "target", "build", ".gradle", ".idea", ".venv", "venv", "env",
        "__pycache__", ".pytest_cache", "site-packages", "dist-info",
        "node_modules",
    }
    for base in roots:
        if not base.exists():
            continue
        if base == ROOT:
            candidates = list(base.glob("docker-compose.imd.yml"))
        else:
            candidates = [p for p in base.rglob("*") if p.is_file()]
        for p in candidates:
            if any(part in excluded_parts for part in p.parts):
                continue
            if p.suffix.lower() in INCLUDE_PATTERNS:
                found.append(p)
    return sorted(set(found), key=lambda p: str(p).lower())


def rel(path):
    return str(path.relative_to(ROOT)).replace("/", "\\")


def clean_annotation_arg(value):
    if not value:
        return ""
    value = re.sub(r"\s+", " ", value)
    return value.replace("\"", "").replace("{", "").replace("}", "").strip()


def parse_controller_routes(path):
    text = read_text(path)
    if "@Controller" not in text and "@RestController" not in text:
        return []
    class_prefix = ""
    class_match = re.search(r"@RequestMapping\s*\((.*?)\)\s*public class", text, re.S)
    if class_match:
        class_prefix = clean_annotation_arg(class_match.group(1))
    routes = []
    for m in re.finditer(r"@(GetMapping|PostMapping|PutMapping|DeleteMapping|RequestMapping)\s*(?:\((.*?)\))?\s*public\s+[\w<>?,\s]+\s+(\w+)\s*\(", text, re.S):
        method = m.group(1).replace("Mapping", "").upper()
        arg = clean_annotation_arg(m.group(2) or "")
        if not arg and method == "REQUEST":
            continue
        route = (class_prefix + "/" + arg).replace("//", "/")
        route = route if route else "/"
        route = route.replace("value = ", "").replace("path = ", "").replace("consumes = multipart/form-data", "")
        routes.append(f"{method} {route} -> {m.group(3)}()")
    return routes


def parse_mapper_ids(path):
    if path.suffix.lower() != ".xml" or "mapper" not in str(path):
        return []
    text = read_text(path)
    return [f"{tag} {idv}" for tag, idv in re.findall(r"<(select|insert|update|delete)\b[^>]*\bid=[\"']([^\"']+)", text)]


def parse_jsp_links(path):
    if path.suffix.lower() != ".jsp":
        return []
    text = read_text(path)
    patterns = [
        r"fetch\([`\"']([^`\"']+)",
        r"action=[\"']([^\"']+)",
        r"href=[\"'](/[^\"']+)",
        r"location\.href\s*=\s*[`\"']([^`\"']+)",
        r"window\.location\s*=\s*[`\"']([^`\"']+)",
    ]
    values = []
    for pat in patterns:
        values.extend(re.findall(pat, text))
    result = []
    for value in values:
        if value and value not in result and not value.startswith("#"):
            result.append(value)
    return result[:25]


def infer_description(path):
    name = path.name
    if name in MANUAL:
        return MANUAL[name]
    if name in JSP_MANUAL:
        return JSP_MANUAL[name]
    if name in DTO_MANUAL:
        return DTO_MANUAL[name]
    r = rel(path)
    if "\\dto\\" in r:
        return "화면 또는 API에서 값을 주고받기 위한 DTO입니다. DB 객체를 그대로 노출하지 않고 필요한 값만 담는 역할입니다."
    if "\\service\\" in r and name.startswith("I"):
        return "서비스 계층 인터페이스입니다. Controller가 구현체에 직접 의존하지 않고 기능 계약을 기준으로 호출하게 합니다."
    if "\\service\\impl\\" in r:
        return "서비스 구현체입니다. Controller에서 받은 요청을 실제 비즈니스 로직, 외부 API, Mapper 호출로 연결합니다."
    if "\\mapper\\" in r and name.endswith(".java"):
        return "MyBatis Mapper 인터페이스입니다. XML SQL과 연결되어 DB 조회/저장/수정/삭제를 수행합니다."
    if "\\config\\" in r:
        return "프로젝트 실행에 필요한 Spring 설정 또는 초기화 코드입니다."
    if "\\controller\\" in r:
        return "사용자 요청을 받는 Controller입니다. 화면 반환 또는 API 응답을 담당합니다."
    if "\\views\\" in r:
        return "JSP 화면 파일입니다. 사용자가 보는 UI이며 버튼, 폼, fetch 요청으로 Controller/API와 연결됩니다."
    if "\\resources\\mapper\\" in r:
        return "MyBatis SQL XML입니다. Mapper 인터페이스 메서드와 실제 SQL을 연결합니다."
    return "프로젝트 실행 또는 기능 구현에 필요한 보조 코드/설정 파일입니다."


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.size = Pt(16 if level == 1 else 13)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121 if level == 1 else 100)
    return p


def add_para(doc, text, size=10.5, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.size = Pt(size)
    r.bold = bold
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r.font.size = Pt(10)


def add_file_item(doc, path):
    rpath = rel(path)
    add_para(doc, rpath, size=10.5, bold=True)
    add_bullet(doc, "역할: " + infer_description(path))
    routes = parse_controller_routes(path)
    if routes:
        add_bullet(doc, "들어오는 경로: " + " / ".join(routes))
    mapper_ids = parse_mapper_ids(path)
    if mapper_ids:
        add_bullet(doc, "SQL ID: " + ", ".join(mapper_ids))
    links = parse_jsp_links(path)
    if links:
        add_bullet(doc, "화면에서 연결되는 경로/API: " + ", ".join(links))


def group_name(path):
    r = rel(path)
    if "\\controller\\api\\" in r:
        return "API Controller"
    if "\\controller\\" in r:
        return "화면 Controller"
    if "\\service\\impl\\" in r:
        return "Service 구현체"
    if "\\service\\" in r:
        return "Service 인터페이스"
    if "\\mapper\\" in r and path.suffix == ".java":
        return "Mapper 인터페이스"
    if "\\resources\\mapper\\" in r:
        return "Mapper XML SQL"
    if "\\dto\\" in r:
        return "DTO"
    if "\\config\\" in r:
        return "Config/초기화"
    if "\\util\\" in r:
        return "Util"
    if "\\views\\" in r:
        return "JSP 화면"
    if "\\document\\" in r or "\\repository\\" in r or "\\entity\\" in r:
        return "Entity/Document/Repository"
    if r.startswith("imd-service") or r.startswith("deploy") or r.startswith("docker-compose"):
        return "IMD/배포"
    return "기타 설정"


def main():
    doc = Document()
    add_heading(doc, "딥페이크 프로젝트 전체 코드 설명서", 1)
    add_para(doc, "현재 프로젝트에 있는 Java, JSP, MyBatis XML, 설정, IMD/배포 파일을 기준으로 빠짐없이 경로와 역할을 정리했습니다.")

    add_heading(doc, "발표용 큰 구조", 1)
    for line in [
        "화면 진입은 JSP용 Controller가 담당합니다. 예: /login, /signup, /news, /community, /history.",
        "비동기/JSON 처리는 controller.api 패키지의 RestController가 담당합니다. 예: /api/v1/auth, /api/v1/verifications, /api/v1/community/posts.",
        "실제 기능 흐름은 Service에서 처리합니다. Controller는 요청을 받고 Service를 호출하는 역할에 집중합니다.",
        "MariaDB 접근은 MyBatis Mapper 인터페이스와 mapper XML이 담당합니다.",
        "MongoDB는 NewsCacheDocument와 NewsCacheRepository를 통해 뉴스 캐시 저장에 사용합니다.",
        "외부 분석은 Reality Defender API 또는 로컬 IMD Python 서비스로 나뉘고, VerifyService가 결과를 하나의 검증 결과로 합칩니다.",
    ]:
        add_bullet(doc, line)

    add_heading(doc, "메인 기능 분석 흐름", 1)
    for title, steps in FEATURE_FLOWS:
        add_heading(doc, title, 2)
        for idx, step in enumerate(steps, 1):
            add_para(doc, f"{idx}. {step}", size=10)

    add_heading(doc, "화면/URL 진입 경로 요약", 1)
    for path in files():
        routes = parse_controller_routes(path)
        if routes:
            add_para(doc, rel(path), bold=True)
            for route in routes:
                add_bullet(doc, route)

    grouped = {}
    for path in files():
        grouped.setdefault(group_name(path), []).append(path)

    add_heading(doc, "전체 파일별 코드 설명", 1)
    order = [
        "Config/초기화", "화면 Controller", "API Controller", "Service 인터페이스", "Service 구현체",
        "Mapper 인터페이스", "Mapper XML SQL", "DTO", "Entity/Document/Repository", "Util",
        "JSP 화면", "기타 설정", "IMD/배포",
    ]
    for group in order:
        if group not in grouped:
            continue
        add_heading(doc, group, 2)
        for path in grouped[group]:
            add_file_item(doc, path)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
