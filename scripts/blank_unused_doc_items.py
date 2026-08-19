from docx import Document
from docx.shared import Pt


INPUT = r"C:\SpringBootWorks\deepfake2\docs\딥페이크_프로젝트_예상질문_답변_보강_코드설명.docx"
OUTPUT = r"C:\SpringBootWorks\deepfake2\docs\딥페이크_프로젝트_예상질문_답변_보강_코드설명_미사용빈칸.docx"


UNUSED_SECTION_HEADINGS = {
    "3. JPA 적용한 학생들 질문",
    "5. JWT 적용한 학생들 질문",
}

SECTION_END_HEADINGS = {
    "4. 스프링 시큐리티 적용한 학생들 질문",
    "5. JWT 적용한 학생들 질문",
    "6. 파일 업로드 질문",
}

UNUSED_QUESTION_STARTS = {
    "4) 로그인해야만 접근 가능한 URL은 무엇인가요?",
    "8) requestMatchers()는 어떤 역할을 하나요?",
    "10) hasRole() 또는 hasAuthority()를 사용했나요?",
}

UNUSED_BULLET_TITLES = {
    "아이디 찾기 이메일 인증번호 발송",
    "아이디 찾기 인증번호 확인",
    "맵 출력 API 호출",
    "경찰서 마커 표시",
}


def clear_paragraph_keep_answer_label(paragraph):
    paragraph.clear()
    run = paragraph.add_run("답변:")
    run.font.name = "맑은 고딕"
    run.font.size = Pt(10.5)


def clear_bullet_body(paragraph, title):
    paragraph.clear()
    r = paragraph.add_run(f"{title}: ")
    r.bold = True
    r.font.name = "맑은 고딕"
    r.font.size = Pt(10.2)


def main():
    doc = Document(INPUT)
    paragraphs = doc.paragraphs

    in_unused_section = False
    blank_next_answer = False

    for p in paragraphs:
        text = p.text.strip()

        if text in UNUSED_SECTION_HEADINGS:
            in_unused_section = True
            blank_next_answer = False
            continue

        if in_unused_section and text in SECTION_END_HEADINGS and text not in UNUSED_SECTION_HEADINGS:
            in_unused_section = False

        if text in SECTION_END_HEADINGS and text not in UNUSED_SECTION_HEADINGS:
            in_unused_section = False

        if in_unused_section and text.startswith("답변:"):
            clear_paragraph_keep_answer_label(p)
            continue

        if text in UNUSED_QUESTION_STARTS:
            blank_next_answer = True
            continue

        if blank_next_answer and text.startswith("답변:"):
            clear_paragraph_keep_answer_label(p)
            blank_next_answer = False
            continue

        for title in UNUSED_BULLET_TITLES:
            if text.startswith(f"{title}:"):
                clear_bullet_body(p, title)
                break

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
