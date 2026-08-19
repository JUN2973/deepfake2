from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "로그인_회원가입_코드_한줄한줄_상세설명.docx"

FILES = [
    ("로그인 JSP", ROOT / "src/main/webapp/WEB-INF/views/login.jsp"),
    ("회원가입 JSP", ROOT / "src/main/webapp/WEB-INF/views/signup.jsp"),
    ("화면 이동 컨트롤러", ROOT / "src/main/java/kopo/poly/controller/AuthPageController.java"),
    ("인증 API 컨트롤러", ROOT / "src/main/java/kopo/poly/controller/api/AuthApiController.java"),
    ("인증 서비스 인터페이스", ROOT / "src/main/java/kopo/poly/service/IAuthService.java"),
    ("인증 서비스 구현체", ROOT / "src/main/java/kopo/poly/service/impl/AuthService.java"),
    ("이메일 인증 서비스", ROOT / "src/main/java/kopo/poly/service/impl/EmailAuthService.java"),
    ("메일 서비스 인터페이스", ROOT / "src/main/java/kopo/poly/service/IMailService.java"),
    ("메일 서비스 구현체", ROOT / "src/main/java/kopo/poly/service/impl/MailService.java"),
    ("구글 OAuth 컨트롤러", ROOT / "src/main/java/kopo/poly/controller/GoogleOAuthController.java"),
    ("구글 OAuth 서비스", ROOT / "src/main/java/kopo/poly/service/impl/GoogleOAuthService.java"),
    ("사용자 Mapper 인터페이스", ROOT / "src/main/java/kopo/poly/mapper/IUserMapper.java"),
    ("이메일 인증 Mapper 인터페이스", ROOT / "src/main/java/kopo/poly/mapper/IEmailAuthMapper.java"),
    ("사용자 SQL Mapper", ROOT / "src/main/resources/mapper/UserMapper.xml"),
    ("이메일 인증 SQL Mapper", ROOT / "src/main/resources/mapper/EmailAuthMapper.xml"),
    ("공통 API 응답 DTO", ROOT / "src/main/java/kopo/poly/dto/ApiResponse.java"),
    ("로그인 요청 DTO", ROOT / "src/main/java/kopo/poly/dto/LoginRequestDTO.java"),
    ("로그인 성공 사용자 DTO", ROOT / "src/main/java/kopo/poly/dto/AuthenticatedUserDTO.java"),
    ("이메일 중복확인 응답 DTO", ROOT / "src/main/java/kopo/poly/dto/EmailAvailabilityDTO.java"),
    ("인증번호 발송 요청 DTO", ROOT / "src/main/java/kopo/poly/dto/SendCodeRequestDTO.java"),
    ("인증번호 확인 요청 DTO", ROOT / "src/main/java/kopo/poly/dto/VerifyCodeRequestDTO.java"),
    ("최종 회원가입 요청 DTO", ROOT / "src/main/java/kopo/poly/dto/SignupCompleteRequestDTO.java"),
    ("회원 저장 DTO", ROOT / "src/main/java/kopo/poly/dto/UserCreateDTO.java"),
    ("이메일 인증 응답 DTO", ROOT / "src/main/java/kopo/poly/dto/EmailAuthResponseDTO.java"),
    ("이메일 인증 Entity", ROOT / "src/main/java/kopo/poly/entity/EmailAuthEntity.java"),
    ("세션 유틸", ROOT / "src/main/java/kopo/poly/util/SessionUtil.java"),
]


def set_run_font(run, name="맑은 고딕", size=9, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, font="맑은 고딕", size=8.5, color="111827", bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, font, size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def clean_code(line):
    return line.rstrip("\n").replace("\t", "    ")


def starts(line, text):
    return line.strip().startswith(text)


def contains(line, text):
    return text in line


def explain_jsp(line):
    s = line.strip()
    if not s:
        return "빈 줄이다. 코드 구역을 나누어 가독성을 높인다."
    if s.startswith("<%@"):
        return "JSP 페이지 지시어다. HTML 응답의 문자 인코딩을 UTF-8로 맞춰 한글이 깨지지 않게 한다."
    if s == "<%":
        return "JSP 스크립틀릿 시작이다. 서버에서 먼저 실행되는 Java 코드를 작성하는 구간이다."
    if "request.getContextPath()" in s:
        return "현재 애플리케이션의 context path를 가져온다. API 호출과 페이지 이동 주소 앞에 붙여 배포 경로 차이를 흡수한다."
    if s == "%>":
        return "JSP 스크립틀릿 종료다. 이후부터 다시 HTML이 브라우저로 내려간다."
    if s.startswith("<!--"):
        return "HTML 주석이다. 브라우저 화면에는 보이지 않고, 코드 설명 목적으로만 남는다."
    if s.startswith("<!DOCTYPE"):
        return "HTML5 문서임을 브라우저에 알린다."
    if s.startswith("<html"):
        return "HTML 문서의 시작이다. `lang=\"ko\"`는 한국어 문서임을 의미한다."
    if s.startswith("</html"):
        return "HTML 문서의 끝이다."
    if s.startswith("<head"):
        return "브라우저에 필요한 메타 정보, 제목, 외부 라이브러리, 스타일을 모으는 영역을 시작한다."
    if s.startswith("</head"):
        return "head 영역을 끝낸다."
    if s.startswith("<meta charset"):
        return "HTML 문서의 문자 인코딩을 UTF-8로 지정한다."
    if "viewport" in s:
        return "모바일과 데스크톱 화면 크기에 맞게 반응형 배율을 설정한다."
    if s.startswith("<title"):
        return "브라우저 탭에 표시될 페이지 제목을 지정한다."
    if "tailwindcss.com" in s:
        return "Tailwind CSS를 불러온다. 화면의 class 속성 스타일들이 이 라이브러리를 사용한다."
    if "lucide" in s and "<script" in s:
        return "lucide 아이콘 라이브러리를 불러온다. `data-lucide` 아이콘을 실제 SVG로 바꿀 때 사용한다."
    if "postcode.v2.js" in s:
        return "다음 주소검색 API를 불러온다. 회원가입 주소 검색 모달에서 사용한다."
    if s.startswith("<link"):
        return "외부 폰트나 리소스를 불러오는 태그다."
    if s.startswith("<style"):
        return "이 JSP 내부에서 직접 적용할 CSS 영역을 시작한다."
    if s.startswith("</style"):
        return "CSS 영역을 끝낸다."
    if s.startswith("body {") or s.startswith("font-family") or s.startswith("background"):
        return "페이지 전체의 폰트나 배경색을 설정하는 CSS 코드다."
    if s.startswith("<body"):
        return "브라우저 화면에 실제로 보이는 본문 영역을 시작한다."
    if s.startswith("</body"):
        return "본문 영역을 끝낸다."
    if "<form" in s and "loginForm" in s:
        return "로그인 폼이다. 아래 JavaScript가 이 ID를 찾아 submit 이벤트를 연결한다."
    if "<form" in s and "signupForm" in s:
        return "회원가입 폼이다. 최종 회원가입 버튼을 누르면 JavaScript submit 이벤트가 실행된다."
    if "id=\"email\"" in s and "<input" in s:
        return "이메일 입력칸이다. 로그인에서는 로그인 ID로, 회원가입에서는 중복확인/인증/최종가입에 모두 사용된다."
    if "id=\"password\"" in s and "<input" in s:
        return "비밀번호 입력칸이다. `type=\"password\"`라 입력값이 화면에서 가려진다."
    if "id=\"confirmPassword\"" in s:
        return "비밀번호 확인 입력칸이다. 최종 회원가입 전에 비밀번호와 같은지 검사한다."
    if "id=\"name\"" in s and "<input" in s:
        return "이름 입력칸이다. 최종 회원가입 payload의 `name` 값으로 서버에 전송된다."
    if "id=\"phone\"" in s:
        return "전화번호 입력칸이다. 아래 JS에서 숫자만 남기고 하이픈 형식으로 자동 변환한다."
    if "id=\"address\"" in s:
        return "주소 입력칸이다. 직접 입력하지 않고 다음 주소검색 모달에서 선택한 주소가 들어간다."
    if "id=\"detailAddress\"" in s:
        return "상세주소 입력칸이다. 주소 선택 후 화면에 표시되어 사용자가 나머지 주소를 입력한다."
    if "id=\"checkEmailButton\"" in s:
        return "이메일 중복확인 버튼이다. 클릭하면 `handleCheckEmail()`이 실행된다."
    if "id=\"sendCodeButton\"" in s:
        return "인증번호 받기 버튼이다. 중복확인이 성공하기 전까지는 disabled 상태다."
    if "id=\"verifyCode\"" in s:
        return "이메일로 받은 6자리 인증번호를 입력하는 칸이다."
    if "id=\"verifyCodeWrap\"" in s:
        return "인증번호 입력 영역이다. 처음에는 숨겨져 있고 인증번호 발송 성공 후 보인다."
    if "id=\"timerText\"" in s or "id=\"timerWrap\"" in s:
        return "인증번호 유효시간 카운트다운을 화면에 표시하는 영역이다."
    if "id=\"errorBox\"" in s or "id=\"errorText\"" in s:
        return "실패 메시지를 표시하기 위한 영역이다. JS의 `showError()`가 값을 넣는다."
    if "id=\"successBox\"" in s or "id=\"successText\"" in s:
        return "성공 메시지를 표시하기 위한 영역이다. JS의 `showSuccess()`가 값을 넣는다."
    if "id=\"submitBtn\"" in s:
        return "폼 제출 버튼이다. 로그인 또는 최종 회원가입 submit 이벤트를 발생시킨다."
    if "oauth2/authorization/google" in s:
        return "구글 로그인 시작 링크다. 클릭하면 `GoogleOAuthController.authorize()`로 이동한다."
    if "find-account" in s:
        return "아이디/비밀번호 찾기 화면으로 이동하는 버튼 또는 주소다."
    if "signup" in s and "goPage" in s:
        return "회원가입 화면으로 이동하는 버튼이다."
    if "login" in s and "goPage" in s:
        return "로그인 화면으로 이동하는 버튼이다."
    if s.startswith("<div") or s.startswith("</div") or s.startswith("<section") or s.startswith("</section") or s.startswith("<main") or s.startswith("</main"):
        return "화면 레이아웃을 구성하는 HTML 태그다. 기능보다는 배치와 묶음 역할을 한다."
    if s.startswith("<label"):
        return "입력칸의 이름표다. `for` 속성으로 어떤 input과 연결되는지 알려준다."
    if s.startswith("<i"):
        return "lucide 아이콘 자리다. 페이지 마지막의 `lucide.createIcons()`가 실제 아이콘으로 렌더링한다."
    if s.startswith("<button"):
        return "사용자가 클릭할 수 있는 버튼을 시작한다. `type`과 `onclick`에 따라 동작이 달라진다."
    if s.startswith("</button"):
        return "버튼 태그를 닫는다."
    if s.startswith("<a"):
        return "링크 태그다. 클릭하면 `href`에 지정된 주소로 이동한다."
    if s.startswith("</a"):
        return "링크 태그를 닫는다."
    if s.startswith("<svg") or s.startswith("<path"):
        return "구글 로그인 버튼의 구글 로고를 그리는 SVG 코드다."
    if s.startswith("<script"):
        return "JavaScript 코드 영역을 시작한다."
    if s.startswith("</script"):
        return "JavaScript 코드 영역을 끝낸다."
    if s.startswith("const state"):
        return "회원가입 화면에서 공유하는 상태 객체를 선언한다. 인증 완료 여부, 타이머, 중복확인 여부 등을 담는다."
    if re.match(r"^\w+:", s):
        key = s.split(":", 1)[0]
        return f"`state` 객체의 `{key}` 값이다. 회원가입 진행 상태를 저장하는 필드다."
    if s.startswith("const "):
        return "JavaScript 상수 변수를 선언한다. DOM 요소나 입력값을 재사용하기 위해 담아둔다."
    if s.startswith("let "):
        return "JavaScript 블록 변수다. 이후 코드에서 값이 바뀔 수 있다."
    if s.startswith("function ") or s.startswith("async function "):
        return "JavaScript 함수를 선언한다. 버튼 클릭이나 폼 제출 시 호출되는 기능 단위다."
    if "addEventListener" in s:
        return "특정 이벤트가 발생했을 때 실행할 함수를 등록한다."
    if "event.preventDefault()" in s:
        return "브라우저의 기본 폼 제출과 페이지 새로고침을 막고 AJAX 방식으로 처리하게 한다."
    if "fetch(" in s:
        if "/api/v1/auth/login" in s:
            return "로그인 API를 호출한다. Java의 `AuthApiController.login()`으로 연결된다."
        if "/signup/check-email" in s:
            return "회원가입 이메일 중복확인 API를 호출한다. Java의 `AuthApiController.checkEmail()`으로 연결된다."
        if "/signup/send-code" in s:
            return "회원가입 인증번호 발송 API를 호출한다. Java의 `AuthApiController.sendCode()`로 연결된다."
        if "/signup/verify-code" in s:
            return "인증번호 확인 API를 호출한다. Java의 `AuthApiController.verifyCode()`로 연결된다."
        if "/signup/complete" in s:
            return "최종 회원가입 API를 호출한다. Java의 `AuthApiController.complete()`로 연결된다."
        return "서버 API를 비동기로 호출한다."
    if "method:" in s:
        return "HTTP 요청 방식을 지정한다. 로그인/회원가입은 서버 상태를 바꾸므로 POST를 사용한다."
    if "Content-Type" in s:
        return "요청 본문이 JSON 형식임을 서버에 알려준다."
    if "JSON.stringify" in s:
        return "JavaScript 객체를 JSON 문자열로 바꿔 서버 요청 본문에 넣는다."
    if ".value" in s and "document.getElementById" in s:
        return "화면 입력칸에서 사용자가 입력한 값을 읽는다."
    if "response.json" in s:
        return "서버 응답 본문을 JSON 객체로 변환한다."
    if "json && json.success" in s:
        return "서버가 성공 응답을 보냈는지 확인한다."
    if "showError" in s:
        return "사용자에게 실패 메시지를 보여준다."
    if "showSuccess" in s:
        return "사용자에게 성공 메시지를 보여준다."
    if "window.location.href" in s or "goPage(" in s:
        return "브라우저를 다른 주소로 이동시킨다."
    if "sessionStorage.setItem" in s:
        return "로그인 성공 후 메인 화면에서 보여줄 토스트 메시지를 브라우저 sessionStorage에 저장한다."
    if "classList.remove" in s:
        return "HTML 요소에서 클래스를 제거한다. 주로 숨김 상태를 해제할 때 사용한다."
    if "classList.add" in s:
        return "HTML 요소에 클래스를 추가한다. 주로 영역을 숨기거나 스타일을 바꿀 때 사용한다."
    if "classList.toggle" in s:
        return "조건에 따라 클래스를 추가하거나 제거한다."
    if "setInterval" in s:
        return "일정 시간마다 반복 실행되는 타이머를 시작한다."
    if "clearInterval" in s:
        return "실행 중인 타이머를 종료한다."
    if "setTimeout" in s:
        return "정해진 시간 뒤에 코드를 한 번 실행한다."
    if "new URLSearchParams" in s:
        return "현재 URL의 query string을 읽는다. 구글 로그인 실패 여부를 확인할 때 사용한다."
    if "new daum.Postcode" in s:
        return "다음 주소검색 위젯 객체를 생성한다."
    if "oncomplete" in s:
        return "사용자가 주소를 선택했을 때 실행되는 콜백 함수다."
    if "data.zonecode" in s or "data.address" in s:
        return "다음 주소검색 API가 돌려준 주소/우편번호 데이터를 사용한다."
    if "replace(/[^\\d]/g" in s:
        return "전화번호 입력값에서 숫자가 아닌 문자를 제거한다."
    if "return" in s:
        return "현재 함수 실행을 끝내고 값을 돌려주거나 흐름을 중단한다."
    if s in ["};", "});", "}", "};", "});"]:
        return "앞에서 시작한 객체, 함수, 조건문 또는 이벤트 처리 블록을 닫는다."
    return "화면 구성 또는 앞뒤 코드의 일부다. 주변 줄과 함께 하나의 HTML/JavaScript 구조를 완성한다."


def explain_java(line, filename):
    s = line.strip()
    if not s:
        return "빈 줄이다. import, 필드, 메서드 같은 코드 구역을 나누어 읽기 쉽게 한다."
    if s.startswith("package "):
        return "이 Java 파일이 속한 패키지를 선언한다."
    if s.startswith("import "):
        return "다른 패키지의 클래스나 인터페이스를 이 파일에서 사용하기 위해 가져온다."
    if s.startswith("/**") or s.startswith("/*") or s.startswith("*") or s.startswith("//"):
        return "주석이다. 코드 실행에는 영향을 주지 않고 역할 설명을 남긴다."
    if s.startswith("@Controller"):
        return "Spring MVC 컨트롤러로 등록한다. 주로 JSP 화면 이름을 반환한다."
    if s.startswith("@RestController"):
        return "JSON API 컨트롤러로 등록한다. 메서드 반환값이 HTTP 응답 본문으로 나간다."
    if s.startswith("@RequestMapping"):
        return "이 컨트롤러의 공통 URL prefix를 지정한다."
    if s.startswith("@GetMapping") or s.startswith("@PostMapping") or s.startswith("@PutMapping") or s.startswith("@DeleteMapping"):
        return "특정 HTTP method와 URL을 이 Java 메서드에 연결한다."
    if s.startswith("@Service"):
        return "Spring Service 빈으로 등록한다. 비즈니스 로직을 담당한다."
    if s.startswith("@Transactional"):
        return "메서드 실행 중 DB 작업을 하나의 트랜잭션으로 묶는다. 실패하면 롤백할 수 있다."
    if s.startswith("@Override"):
        return "인터페이스나 부모 클래스의 메서드를 구현/재정의한다는 표시다."
    if s.startswith("@Value"):
        return "application.properties 또는 환경변수 값을 필드에 주입한다."
    if s.startswith("@Slf4j") or s.startswith("@RequiredArgsConstructor") or s.startswith("@FunctionalInterface"):
        return "Lombok 또는 Java 기능을 위한 어노테이션이다. 로그 필드, 생성자, 함수형 인터페이스 등을 도와준다."
    if "class " in s and s.startswith("public"):
        return "공개 Java 클래스를 선언한다. 이 파일의 주요 기능 단위다."
    if s.startswith("public interface"):
        return "인터페이스를 선언한다. 구현 클래스가 제공해야 할 메서드 목록을 정의한다."
    if s.startswith("private static final Logger"):
        return "로그를 남기기 위한 Logger 객체를 선언한다."
    if "DATASOURCE_ERROR" in s:
        return "DB 연결 오류가 났을 때 프론트에 내려줄 공통 메시지를 상수로 관리한다."
    if "PasswordEncoder" in s or "BCryptPasswordEncoder" in s:
        return "비밀번호를 BCrypt 방식으로 암호화하거나 비교하기 위한 객체다."
    if s.startswith("private final"):
        return "생성자 주입으로 받을 의존성 필드다. Controller나 Service가 다른 계층을 호출할 때 사용한다."
    if "public AuthApiController" in s or "public AuthService" in s or "public GoogleOAuthController" in s or "public GoogleOAuthService" in s:
        return "생성자다. 필요한 의존성을 주입받아 필드에 저장한다."
    if "this." in s:
        return "생성자로 받은 값을 현재 객체의 필드에 저장한다."
    if "return \"login\"" in s:
        return "`login.jsp` 화면을 렌더링하라는 의미로 view 이름 `login`을 반환한다."
    if "return \"signup\"" in s:
        return "`signup.jsp` 화면을 렌더링하라는 의미로 view 이름 `signup`을 반환한다."
    if "return \"find-account\"" in s:
        return "계정 찾기 JSP 화면으로 이동하기 위한 view 이름을 반환한다."
    if "return \"mypage\"" in s:
        return "마이페이지 JSP 화면으로 이동하기 위한 view 이름을 반환한다."
    if "handleAuthRequest" in s and "return" in s:
        return "인증 API 로직을 공통 예외 처리 메서드로 감싼다. 성공/실패 응답 형식을 통일한다."
    if "authService.login" in s:
        return "로그인 검증을 Service 계층에 위임한다. 성공하면 사용자 DTO를 받는다."
    if "request.changeSessionId" in s:
        return "로그인 성공 직후 세션 ID를 교체한다. 세션 고정 공격을 줄이기 위한 보안 처리다."
    if "session.setAttribute" in s:
        return "로그인 사용자 정보를 서버 세션에 저장한다. 이후 화면과 API에서 로그인 상태 판단에 사용한다."
    if "ApiResponse.ok" in s:
        return "성공 응답 JSON을 만든다."
    if "ApiResponse.fail" in s:
        return "실패 응답 JSON을 만든다."
    if "session.invalidate" in s:
        return "현재 세션을 무효화한다. 로그아웃 또는 회원탈퇴 후 로그인 상태를 제거한다."
    if "SessionUtil.getUserId" in s:
        return "세션에서 현재 로그인 사용자의 USER_ID를 꺼낸다."
    if "authService.checkEmail" in s:
        return "회원가입 이메일 중복확인 로직을 Service에 위임한다."
    if "authService.sendSignupCode" in s:
        return "회원가입 인증번호 발송 로직을 Service에 위임한다."
    if "authService.verifySignupCode" in s:
        return "회원가입 인증번호 확인 로직을 Service에 위임한다."
    if "authService.completeSignup" in s:
        return "최종 회원가입 저장 로직을 Service에 위임한다."
    if "catch (AuthServiceException" in s:
        return "서비스에서 발생한 인증/검증 실패 예외를 잡는다."
    if "catch (DataAccessException" in s:
        return "DB 접근 중 발생한 예외를 잡는다."
    if "catch (Exception" in s:
        return "예상하지 못한 일반 예외를 잡는다."
    if "log.error" in s or "log.warn" in s:
        return "서버 로그에 오류 또는 경고 내용을 남긴다."
    if "LoginRequestDTO" in s:
        return "로그인 요청의 이메일과 비밀번호를 담는 DTO와 관련된 줄이다."
    if "SignupCompleteRequestDTO" in s:
        return "최종 회원가입 요청 데이터를 담는 DTO와 관련된 줄이다."
    if "SendCodeRequestDTO" in s:
        return "이메일 중복확인/인증번호 발송 요청의 이메일 값을 담는 DTO와 관련된 줄이다."
    if "VerifyCodeRequestDTO" in s:
        return "인증번호 확인 요청의 이메일과 코드를 담는 DTO와 관련된 줄이다."
    if "AuthenticatedUserDTO" in s:
        return "로그인 성공 후 세션 저장에 필요한 사용자 id, 이름, 이메일을 담는 DTO와 관련된 줄이다."
    if "String email =" in s:
        return "요청에서 이메일을 꺼내 null 방어와 공백 제거, 소문자 변환 등을 수행한다."
    if "String password =" in s:
        return "요청에서 비밀번호를 꺼낸다. 비밀번호는 공백 제거 없이 사용자가 입력한 원문을 비교한다."
    if "isBlank()" in s:
        return "문자열이 비어 있거나 공백뿐인지 검사한다."
    if "throw new AuthServiceException" in s:
        return "검증 실패를 의미하는 사용자 정의 예외를 발생시킨다. 컨트롤러에서 JSON 실패 응답으로 바뀐다."
    if "userMapper.selectUserByEmail" in s:
        return "DB의 app_user 테이블에서 이메일로 사용자를 조회한다. UserMapper.xml의 selectUserByEmail SQL과 연결된다."
    if "userMapper.existsByEmail" in s:
        return "DB의 app_user 테이블에서 같은 이메일이 이미 있는지 확인한다."
    if "userMapper.insertUser" in s:
        return "최종 회원가입 정보를 app_user 테이블에 저장한다."
    if "userMapper.updatePasswordByEmail" in s:
        return "이메일 기준으로 비밀번호 해시를 갱신한다."
    if "passwordMatches" in s:
        return "입력 비밀번호와 저장된 비밀번호 해시가 일치하는지 비교한다."
    if "isLegacySha256" in s:
        return "예전 SHA-256 해시 형식인지 확인한다."
    if "HashUtil.sha256" in s:
        return "입력 비밀번호를 SHA-256으로 해시해서 레거시 저장값과 비교한다."
    if "PASSWORD_ENCODER.encode" in s:
        return "비밀번호를 BCrypt 해시로 변환한다. DB에는 평문 비밀번호를 저장하지 않는다."
    if "PASSWORD_ENCODER.matches" in s:
        return "입력 비밀번호와 BCrypt 해시가 일치하는지 검사한다."
    if "emailAuthService.sendAuthCode" in s:
        return "이메일 인증번호 생성, DB 저장, 메일 발송을 EmailAuthService에 맡긴다."
    if "emailAuthService.verifyAuthCode" in s:
        return "사용자가 입력한 인증번호가 DB의 최신 인증번호와 맞는지 EmailAuthService에 맡긴다."
    if "emailAuthService.isEmailVerified" in s:
        return "최종 가입 또는 비밀번호 재설정 전에 이메일 인증 완료 여부를 DB 기준으로 확인한다."
    if "emailAuthService.expireEmailVerification" in s:
        return "사용이 끝난 이메일 인증 상태를 만료 처리해 재사용을 막는다."
    if "UserCreateDTO" in s:
        return "DB에 저장할 회원 정보를 담는 DTO를 만들거나 사용하는 줄이다."
    if "buildFullAddress" in s:
        return "우편번호, 기본주소, 상세주소를 하나의 문자열로 합치는 helper와 관련된 줄이다."
    if "SecureRandom" in s:
        return "예측하기 어려운 난수 생성에 사용한다. 인증번호나 OAuth state 생성에 필요하다."
    if "MAX_SENDS_PER_WINDOW" in s:
        return "일정 시간 동안 허용할 인증번호 발송 요청 횟수 제한값이다."
    if "MAX_VERIFY_FAILURES_PER_WINDOW" in s:
        return "일정 시간 동안 허용할 인증번호 검증 실패 횟수 제한값이다."
    if "RATE_LIMIT_WINDOW" in s:
        return "요청/실패 횟수 제한을 계산할 시간 창이다."
    if "generateAuthCode" in s:
        return "6자리 이메일 인증번호를 생성하는 helper를 호출하거나 정의한다."
    if "LocalDateTime.now().plusMinutes(5)" in s:
        return "인증번호 만료 시간을 현재 시각 기준 5분 뒤로 설정한다."
    if "expireAllByEmail" in s:
        return "같은 이메일의 기존 인증번호를 만료 처리한다."
    if "insertEmailAuth" in s:
        return "새 이메일 인증번호 기록을 email_auth 테이블에 저장한다."
    if "selectLatestByEmail" in s:
        return "해당 이메일의 최신 인증번호 기록을 email_auth 테이블에서 조회한다."
    if "markVerified" in s:
        return "인증번호가 맞을 때 email_auth의 verified 값을 Y로 변경한다."
    if "mailService.doSendMail" in s:
        return "구성한 메일 DTO를 실제 메일 발송 서비스로 보낸다."
    if "MailDTO" in s:
        return "메일 수신자, 제목, 내용을 담는 DTO와 관련된 줄이다."
    if "RestClient" in s:
        return "구글 OAuth 서버와 HTTP 통신할 때 사용하는 Spring HTTP 클라이언트다."
    if "GOOGLE_AUTH_URL" in s or "GOOGLE_TOKEN_URL" in s or "GOOGLE_USERINFO_URL" in s:
        return "구글 OAuth 인증, 토큰 교환, 사용자 정보 조회 엔드포인트 상수다."
    if "hasGoogleConfig" in s:
        return "구글 OAuth client id/secret 설정이 유효한지 확인한다."
    if "createState" in s:
        return "OAuth CSRF 방지를 위한 state 난수를 생성한다."
    if "buildAuthorizationUrl" in s:
        return "구글 로그인 화면으로 보낼 인증 URL을 만든다."
    if "loginWithCode" in s:
        return "구글 콜백 code를 access token으로 교환하고 서비스 사용자로 로그인시킨다."
    if "requestToken" in s:
        return "구글 토큰 엔드포인트에 code를 보내 access token을 받는다."
    if "requestProfile" in s:
        return "구글 userinfo 엔드포인트에서 사용자 프로필을 조회한다."
    if "findOrCreateUser" in s:
        return "구글 사용자와 서비스 회원을 연결하거나 신규 소셜 회원을 만든다."
    if "selectUserByOauth" in s:
        return "OAuth 제공자와 제공자 ID로 기존 소셜 로그인 사용자를 찾는다."
    if "updateOauthByEmail" in s:
        return "동일 이메일의 기존 계정에 구글 OAuth 정보를 연결한다."
    if "insertSocialUser" in s:
        return "처음 방문한 구글 사용자를 비밀번호 없는 소셜 회원으로 저장한다."
    if s.startswith("return "):
        return "메서드 실행 결과를 호출한 곳으로 반환한다."
    if s in ["}", "};", ");", "});"]:
        return "앞에서 시작한 클래스, 메서드, 조건문, 람다, 객체 생성 블록 등을 닫는다."
    return "Java 코드의 한 줄이다. 주변 줄과 함께 클래스, 메서드, 조건문, 객체 생성 또는 예외 처리 구조를 완성한다."


def explain_xml(line):
    s = line.strip()
    if not s:
        return "빈 줄이다. SQL 구역을 나누어 읽기 쉽게 한다."
    if s.startswith("<?xml"):
        return "XML 파일 선언이다. 인코딩을 UTF-8로 지정한다."
    if s.startswith("<!--"):
        return "XML 주석이다. SQL 실행에는 영향을 주지 않는다."
    if s.startswith("<!DOCTYPE"):
        return "MyBatis mapper XML 문서 형식을 지정한다."
    if s.startswith("<mapper"):
        return "이 XML을 특정 Java Mapper 인터페이스와 연결한다."
    if s.startswith("</mapper"):
        return "MyBatis mapper XML을 닫는다."
    if s.startswith("<select"):
        return "SELECT SQL 매핑을 시작한다. Java Mapper의 같은 id 메서드와 연결된다."
    if s.startswith("</select"):
        return "SELECT SQL 매핑을 끝낸다."
    if s.startswith("<insert"):
        return "INSERT SQL 매핑을 시작한다. DB에 새 행을 저장할 때 사용한다."
    if s.startswith("</insert"):
        return "INSERT SQL 매핑을 끝낸다."
    if s.startswith("<update"):
        return "UPDATE SQL 매핑을 시작한다. DB의 기존 행을 수정할 때 사용한다."
    if s.startswith("</update"):
        return "UPDATE SQL 매핑을 끝낸다."
    if s.startswith("<delete"):
        return "DELETE SQL 매핑을 시작한다. DB 행 삭제에 사용한다."
    if s.startswith("</delete"):
        return "DELETE SQL 매핑을 끝낸다."
    if s.upper().startswith("SELECT"):
        return "DB에서 데이터를 조회하는 SQL이다."
    if s.upper().startswith("INSERT"):
        return "DB 테이블에 새 데이터를 저장하는 SQL이다."
    if s.upper().startswith("UPDATE"):
        return "DB 테이블의 기존 데이터를 수정하는 SQL이다."
    if s.upper().startswith("DELETE"):
        return "DB 테이블의 데이터를 삭제하는 SQL이다."
    if s.upper().startswith("FROM"):
        return "SQL이 대상으로 삼는 테이블을 지정한다."
    if s.upper().startswith("WHERE"):
        return "SQL 실행 대상을 조건으로 제한한다."
    if "COUNT(1)" in s:
        return "조건에 맞는 행 개수를 센다. 이메일 중복확인에 사용된다."
    if "#{email}" in s:
        return "Java Mapper 메서드에서 전달한 email 파라미터가 들어가는 자리다."
    if "#{passwordHash}" in s:
        return "서비스에서 BCrypt로 만든 비밀번호 해시가 들어가는 자리다."
    if "ORDER BY" in s:
        return "조회 결과의 정렬 기준을 지정한다. 최신 인증번호를 찾기 위해 id 내림차순을 사용한다."
    if "LIMIT 1" in s:
        return "조회 결과를 1개로 제한한다."
    return "SQL 문장의 일부다. 앞뒤 줄과 함께 조회, 저장, 수정, 삭제 쿼리를 완성한다."


def explain_line(path, line):
    suffix = path.suffix.lower()
    if suffix == ".jsp":
        return explain_jsp(line)
    if suffix == ".xml":
        return explain_xml(line)
    return explain_java(line, path.name)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9)
    return doc


def add_intro(doc):
    p = doc.add_paragraph()
    r = p.add_run("로그인 / 회원가입 코드 한 줄 한 줄 상세 설명")
    set_run_font(r, size=18, bold=True, color="1F4D78")
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    r = p.add_run("범위: login.jsp, signup.jsp에서 시작해 Controller, Service, DTO, Mapper XML, DB 연결, 구글 로그인까지 이어지는 관련 코드")
    set_run_font(r, size=9.5, color="475569")
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    r = p.add_run("읽는 방법: 각 파일별로 줄 번호, 실제 코드, 해당 줄의 의미를 함께 적었다. fetch URL, service 호출, mapper 호출처럼 다음 파일로 이동하는 줄은 연결 지점을 설명에 포함했다.")
    set_run_font(r, size=9.5, color="111827")
    p.paragraph_format.space_after = Pt(10)


def add_file_table(doc, title, path):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(5)
    r = h.add_run(f"{title} - {path.relative_to(ROOT).as_posix()}")
    set_run_font(r, size=13, bold=True, color="2E74B5")

    lines = path.read_text(encoding="utf-8").splitlines()
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    widths = [Inches(0.55), Inches(5.25), Inches(4.0)]
    headers = ["줄", "코드", "설명"]
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_text(cell, headers[idx], size=8.5, bold=True, color="0F172A")
        set_cell_shading(cell, "E8EEF5")
        cell.width = widths[idx]

    for no, line in enumerate(lines, 1):
        row = table.add_row()
        code = clean_code(line)
        explanation = explain_line(path, line)

        set_cell_text(row.cells[0], str(no), size=8, color="334155")
        set_cell_text(row.cells[1], code if code else " ", font="Consolas", size=7.3, color="111827")
        set_cell_text(row.cells[2], explanation, size=8.2, color="111827")
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]


def main():
    doc = setup_doc()
    add_intro(doc)
    for title, path in FILES:
        if path.exists():
            add_file_table(doc, title, path)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
