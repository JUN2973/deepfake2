from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches


ROOT = Path(r"C:\SpringBootWorks\deepfake2")
DOCS = ROOT / "docs"
QNA_OUT = DOCS / "딥페이크_프로젝트_예상질문_답변_보강_코드설명_수정본.docx"
FULL_OUT = DOCS / "딥페이크_프로젝트_전체코드_경로_메인기능_설명서_수정본.docx"


FONT = "Malgun Gothic"
BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(30, 41, 59)
GRAY = RGBColor(71, 85, 105)


def set_doc_style(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"].font.size = Pt(10.2)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    styles["Normal"].paragraph_format.space_after = Pt(4)
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, DARK),
        ("Heading 3", 11.5, GRAY),
    ]:
        styles[name].font.name = FONT
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = color
        styles[name].paragraph_format.space_before = Pt(8)
        styles[name].paragraph_format.space_after = Pt(5)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.name = FONT
    r.font.size = Pt(20)
    r.bold = True
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.name = FONT
    r.font.size = Pt(10.5)
    r.font.color.rgb = GRAY


def h(doc, text, level=1):
    doc.add_paragraph(text, style=f"Heading {level}")


def p(doc, text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(10.2)
    run.bold = bold
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(10.0)
    return para


def numbered(doc, text):
    para = doc.add_paragraph(style="List Number")
    run = para.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(10.0)
    return para


def qa(doc, question, answer):
    p(doc, question, bold=True)
    p(doc, "답변: " + answer)


def build_qna():
    doc = Document()
    set_doc_style(doc)
    add_title(
        doc,
        "딥페이크 프로젝트 예상 질문·답변 정리본",
        "현재 C:\\SpringBootWorks\\deepfake2 코드 기준으로 발표 답변과 코드 설명을 다시 맞춘 수정본",
    )

    h(doc, "1. 프로젝트 기본 구조", 1)
    qa(doc, "1) 본인 프로젝트의 전체 패키지 구조를 설명해보세요.",
       "kopo.poly 패키지를 기준으로 config, controller, controller.api, service, service.impl, mapper, dto, document, repository, util로 나눴습니다. "
       "JSP 화면 이동은 controller 패키지가 맡고, JSON API는 controller.api가 맡습니다. 실제 기능 흐름은 Service가 처리하고, MariaDB는 MyBatis Mapper와 XML, MongoDB 뉴스 캐시는 NewsCacheRepository가 담당합니다.")
    qa(doc, "2) 현재 빌드 환경은 무엇인가요?",
       "Maven이 아니라 Gradle 기반입니다. build.gradle에서 Spring Boot 3.3.12, Java 17 toolchain, war 패키징, MyBatis, Spring MVC, Mail, MongoDB, AWS S3 SDK, Spring Security Crypto, JSP/JSTL, jsoup 의존성을 사용합니다.")
    qa(doc, "3) 서버 시작 클래스는 무엇인가요?",
       "Deepfake2Application입니다. @SpringBootApplication으로 자동 설정과 컴포넌트 스캔을 수행하고, @MapperScan(\"kopo.poly.mapper\")로 MyBatis Mapper 인터페이스를 스캔합니다.")
    qa(doc, "4) Controller, Service, Mapper의 역할 차이는 무엇인가요?",
       "Controller는 요청을 받고 화면 이름이나 JSON 응답을 반환합니다. Service는 파일 검증, 외부 API 호출, 회원 인증, 뉴스 캐시, 커뮤니티 권한 확인 같은 비즈니스 흐름을 처리합니다. Mapper는 MyBatis XML SQL과 연결되어 MariaDB에 접근합니다.")
    qa(doc, "5) 현재 주요 화면 JSP는 무엇인가요?",
       "메인은 home.jsp, 로그인은 login.jsp, 회원가입은 signup.jsp, 계정 찾기는 find-account.jsp, 마이페이지는 mypage.jsp, 커뮤니티는 community.jsp/create-post.jsp/post-detail.jsp, 뉴스는 news.jsp/news-detail.jsp입니다. 검증 기록과 결과 쪽은 현재 HistoryController가 deepfake-history.jsp, ResultController가 deepfake-result.jsp를 반환합니다.")

    h(doc, "2. 딥페이크 검증 기능", 1)
    qa(doc, "6) 가장 중요한 기능 하나를 요청 흐름으로 설명해보세요.",
       "이미지 딥페이크 검증 기능입니다. 사용자가 파일을 업로드하면 UploadController 또는 VerificationApiController가 file 파라미터를 받고 VerifyService.createVerification(file, userId)를 호출합니다. VerifyService는 파일 형식과 용량을 검사하고, 로컬 uploads 또는 S3에 저장한 뒤 Reality Defender 또는 IMD 분석 클라이언트를 호출합니다. 결과는 verification 테이블에 저장되고 상세 화면에서 다시 조회됩니다.")
    qa(doc, "7) 파일 검증은 어디서 하나요?",
       "VerifyService.validateImageOnly()에서 처리합니다. contentType이 image/로 시작하는지, 확장자가 png/jpg/jpeg/bmp/webp/heic/heif인지, 파일 크기가 32MB 이하인지 확인합니다.")
    qa(doc, "8) 분석 전에 사전 필터링을 하는 이유는 무엇인가요?",
       "너무 작은 이미지, 흐린 이미지, 얼굴 분석에 부적절한 이미지, 사람이 너무 많아 단일 얼굴 분석이 어려운 이미지는 외부 API 결과 신뢰도가 낮아질 수 있습니다. 그래서 ImageIO 기반으로 크기, 흐림, 피부색 영역 등을 먼저 확인하고 부적절하면 NOT_APPLICABLE 결과를 만듭니다.")
    qa(doc, "9) Reality Defender와 IMD는 각각 어떤 역할인가요?",
       "Reality DefenderClient는 실제 딥페이크 판정 API를 호출합니다. IMD는 로컬 Python FastAPI 서비스로, deepfake.client.mode=imd일 때 주 분석기로 쓸 수 있고, real 모드에서는 ImdHeatmapClient가 보조 히트맵과 의심 영역 시각화에 사용됩니다.")
    qa(doc, "10) 분석 결과는 DB에 어떤 값으로 저장되나요?",
       "verification 테이블에 원본 파일명, MIME 타입, 파일 크기, objectKey, publicUrl, verdict, score, apiProvider, apiRaw, userId를 저장합니다. apiRaw에는 외부 분석 원본 JSON 또는 병합 JSON을 저장해서 상세 화면에서 근거와 히트맵을 다시 만들 수 있게 했습니다.")
    qa(doc, "11) 파일 저장소는 어떻게 바꿀 수 있나요?",
       "IObjectStorageService 인터페이스를 두고 local 모드에서는 DummyObjectStorageService가 uploads 폴더에 저장하고, s3 모드에서는 S3ObjectStorageService가 Amazon S3에 업로드합니다. application.properties의 app.storage.type=${APP_STORAGE_TYPE:local} 설정으로 전환합니다.")
    qa(doc, "12) S3 배포 시 credentials 오류가 나는 이유와 해결 방법은 무엇인가요?",
       "APP_STORAGE_TYPE=s3인데 AWS_ACCESS_KEY_ID/AWS_SECRET_ACCESS_KEY 또는 EC2 IAM Role이 없으면 AWS SDK가 자격증명을 찾지 못해 업로드가 실패합니다. 로컬 개발은 APP_STORAGE_TYPE=local로 두고, 운영 S3 배포는 IAM Role을 붙이거나 AWS 키 환경변수를 넣어야 합니다.")
    qa(doc, "13) 결과 상세 화면의 XSS 방어는 어디서 하나요?",
       "ResultController에서 apiRaw와 analysisJson을 JSP script 안에 넣기 전에 safeJsonForScript()로 <, >, &, U+2028, U+2029 문자를 이스케이프합니다. JSON이 script 태그를 깨고 실행되는 것을 막기 위한 처리입니다.")

    h(doc, "3. 회원, 인증, 보안", 1)
    qa(doc, "14) 로그인 처리는 어떻게 동작하나요?",
       "AuthApiController.login()이 email/password를 받고 IUserMapper.selectUserByEmail()로 회원을 찾습니다. 비밀번호는 BCryptPasswordEncoder로 검증하며, 과거 SHA-256 해시가 남아 있으면 로그인 성공 시 BCrypt로 재저장합니다. 성공 후 request.changeSessionId()로 세션 고정 공격을 줄이고 USER_ID, USER_NAME, USER_EMAIL을 세션에 저장합니다.")
    qa(doc, "15) 비밀번호 저장 방식은 무엇인가요?",
       "현재 신규 회원가입과 비밀번호 변경/재설정은 BCrypt로 저장합니다. 기존 SHA-256 64자리 해시는 legacy 값으로만 인식하고, 로그인 성공 시 BCrypt로 마이그레이션합니다.")
    qa(doc, "16) 회원가입 흐름을 설명해보세요.",
       "check-email로 중복 이메일을 확인하고, send-code로 이메일 인증번호를 발송한 뒤 verify-code로 인증을 완료합니다. 마지막 complete에서 이메일 인증 여부, 비밀번호 길이, 이름, 전화번호, 주소를 검증하고 app_user 테이블에 회원을 저장합니다.")
    qa(doc, "17) Google OAuth는 Spring Security 설정으로 처리하나요?",
       "아니요. 현재는 SecurityConfig나 GoogleOAuth2SuccessHandler 방식이 아니라 GoogleOAuthController가 직접 OAuth 흐름을 처리합니다. /oauth2/authorization/google에서 state를 만들고 Google 인증 URL로 이동시키며, /login/oauth2/code/google 콜백에서 토큰과 사용자 정보를 받아 회원 연결 또는 소셜 회원 생성을 수행합니다.")
    qa(doc, "18) Google OAuth에서 state 값은 왜 사용하나요?",
       "CSRF 성격의 OAuth 위조 요청을 막기 위해 사용합니다. authorize 단계에서 SecureRandom으로 state를 만들고 세션에 저장한 뒤, callback 단계에서 Google이 돌려준 state와 세션의 state가 일치하는지 확인합니다.")
    qa(doc, "19) OriginProtectionFilter는 무엇을 막나요?",
       "POST, PUT, PATCH, DELETE 같은 상태 변경 요청에서 Origin 또는 Referer가 현재 서버 origin 또는 app.security.allowed-origins에 포함된 값과 다르면 403을 반환합니다. CORS만으로 막기 어려운 외부 사이트의 상태 변경 요청을 한 번 더 방어하는 필터입니다.")
    qa(doc, "20) 세션 권한 확인은 어디서 하나요?",
       "SessionUtil.getUserId()와 SessionUtil.canAccessVerification()을 사용합니다. 검증 결과 상세 조회는 세션 userId와 기록의 userId가 일치하거나, 비로그인 업로드 직후 세션에 기억한 검증 ID인 경우에만 허용합니다.")

    h(doc, "4. 외부 API와 설정", 1)
    qa(doc, "21) 프로젝트에서 사용하는 외부 API는 무엇인가요?",
       "Reality Defender API, Naver News API, Google OAuth2 API를 사용합니다. 추가로 로컬 Python IMD FastAPI 서비스를 HTTP로 호출해 분석 또는 히트맵 시각화를 수행합니다.")
    qa(doc, "22) 외부 API 호출 방식은 무엇인가요?",
       "현재 공통 HTTP 클라이언트는 RestTemplate이 아니라 Spring의 RestClient입니다. AppConfig에서 RestClient Bean을 만들고 connect/read timeout을 설정합니다. GoogleOAuthController와 RealityDefenderClient도 RestClient를 사용합니다.")
    qa(doc, "23) 타임아웃 설정은 어디에 있나요?",
       "application.properties의 app.http.connect-timeout-ms, app.http.read-timeout-ms로 공통 RestClient 타임아웃을 설정합니다. IMD 히트맵 요청은 imd.request-timeout-ms로 별도 긴 read timeout을 사용합니다.")
    qa(doc, "24) API Key와 민감정보는 어디에 저장하나요?",
       "코드에 직접 쓰지 않고 application.properties가 환경변수 또는 .env.properties 값을 읽도록 했습니다. 예: RD_API_KEY, NAVER_CLIENT_ID, NAVER_CLIENT_SECRET, GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET, DB_URL, MONGODB_URI 등입니다. .env.properties는 Git에 올리지 않는 전제입니다.")
    qa(doc, "25) Naver News API 결과는 어떻게 캐시하나요?",
       "NewsService가 로컬 메모리 캐시와 MongoDB news_cache를 함께 사용합니다. NewsCacheDocument는 뉴스 제목, 설명, 링크, 발행일, 제공처, 태그, 캐시 만료 정보 등을 담고 NewsCacheRepository가 MongoRepository로 저장/조회합니다.")

    h(doc, "5. DB와 MyBatis", 1)
    qa(doc, "26) MariaDB에는 어떤 주요 테이블이 있나요?",
       "app_user, email_auth, verification, verification_record, community_post, community_comment, community_post_like, community_comment_like 등이 있습니다. 초기 스키마 보정은 AppUserSchemaInitializer, EmailAuthSchemaInitializer, VerificationSchemaInitializer, DetectionSchemaInitializer가 담당합니다.")
    qa(doc, "27) MyBatis Mapper와 XML은 어떻게 연결되나요?",
       "Deepfake2Application의 @MapperScan이 kopo.poly.mapper 패키지를 스캔하고, application.properties의 mybatis.mapper-locations=classpath:/mapper/*.xml 설정으로 XML SQL을 연결합니다. Mapper 인터페이스의 메서드명과 XML의 id가 매칭됩니다.")
    qa(doc, "28) JPA Entity를 주로 사용하지 않는 이유는 무엇인가요?",
       "업무 데이터는 MyBatis와 SQL XML 중심으로 처리했습니다. 다만 MongoDB 뉴스 캐시는 NewsCacheDocument와 NewsCacheRepository를 사용해 Spring Data MongoDB 방식으로 처리합니다.")

    h(doc, "6. 커뮤니티, 뉴스, 기타 화면", 1)
    qa(doc, "29) 커뮤니티 글 작성 권한은 어떻게 확인하나요?",
       "CommunityApiController가 세션의 USER_ID와 USER_NAME을 확인합니다. 수정/삭제는 CommunityService에서 게시글 또는 댓글의 userId와 현재 세션 userId를 비교해 작성자만 가능하게 합니다.")
    qa(doc, "30) 댓글과 대댓글은 어떻게 구성하나요?",
       "community_comment 테이블에 parent_id를 두고, CommunityService.buildCommentTree()가 parentId 기준으로 댓글과 대댓글 구조를 만듭니다.")
    qa(doc, "31) 뉴스 상세는 어떻게 보여주나요?",
       "NewsController의 /news/{id}가 NewsService.getNewsDetailView() 계열 흐름으로 MongoDB 또는 로컬 캐시에서 문서를 찾고 NewsViewDTO로 변환해 news-detail.jsp에 전달합니다.")
    qa(doc, "32) FAQ와 신고 화면은 어떤 구조인가요?",
       "FaqController는 /faq에서 faq.jsp를 반환하고, ReportController는 /report에서 report.jsp를 반환합니다. 주로 정적 안내와 화면 UI 중심입니다.")

    h(doc, "7. 배포와 운영 답변", 1)
    qa(doc, "33) 배포 산출물은 무엇인가요?",
       "Gradle bootWar가 deepfake2-0.0.1-SNAPSHOT.war를 생성합니다. deploy 폴더에는 AWS 배포 가이드, IMD Docker 배포 가이드, 운영 docker-compose.imd.prod.yml, 환경변수 예시가 있습니다.")
    qa(doc, "34) 로컬과 운영의 설정 차이는 무엇인가요?",
       "로컬은 APP_STORAGE_TYPE=local로 uploads 폴더를 쓰는 것이 편하고, 운영은 APP_STORAGE_TYPE=s3로 S3를 사용할 수 있습니다. 운영에서는 DB, MongoDB, Mail, Reality Defender, Naver, Google OAuth, S3 값을 환경변수 또는 배포용 .env.properties로 주입해야 합니다.")
    qa(doc, "35) 발표 때 보안 개선점으로 무엇을 말할 수 있나요?",
       "BCrypt 적용, legacy SHA-256 자동 마이그레이션, 로그인/OAuth 성공 시 changeSessionId(), OriginProtectionFilter의 상태 변경 요청 origin 검사, 상세 결과 JSON script 이스케이프, 세션 기반 검증 기록 접근 제어를 말할 수 있습니다.")

    doc.save(QNA_OUT)


def file_section(doc, title, items):
    h(doc, title, 2)
    for path, role in items:
        p(doc, path, bold=True)
        bullet(doc, role)


def build_full():
    doc = Document()
    set_doc_style(doc)
    add_title(
        doc,
        "딥페이크 프로젝트 전체 코드 설명서",
        "현재 코드 기준 경로, 역할, 요청 흐름, 주요 설정을 발표용으로 정리한 수정본",
    )

    h(doc, "발표용 큰 구조", 1)
    for line in [
        "JSP 화면 진입은 controller 패키지가 담당합니다. 예: /, /login, /signup, /news, /community, /history.",
        "비동기/JSON 처리는 controller.api 패키지의 RestController가 담당합니다. 예: /api/v1/auth, /api/v1/verifications, /api/v1/community/posts.",
        "실제 기능 흐름은 Service가 처리하고, MariaDB 접근은 MyBatis Mapper 인터페이스와 mapper XML이 담당합니다.",
        "뉴스 캐시는 MongoDB의 news_cache 컬렉션에 NewsCacheDocument 형태로 저장합니다.",
        "이미지 분석은 Reality Defender API와 로컬 IMD Python 서비스를 조합하고, VerifyService가 결과를 하나의 검증 결과로 합칩니다.",
    ]:
        bullet(doc, line)

    h(doc, "메인 딥페이크 분석 기능 흐름", 1)
    for step in [
        "사용자는 / 또는 /home에서 home.jsp를 통해 이미지를 선택합니다.",
        "업로드는 /upload 또는 /api/v1/verifications로 multipart/form-data 요청을 보냅니다.",
        "UploadController 또는 VerificationApiController가 file을 꺼내 VerifyService.createVerification(file, userId)를 호출합니다.",
        "VerifyService.validateImageOnly()가 MIME 타입, 확장자, 32MB 이하 용량을 검사합니다.",
        "IObjectStorageService를 통해 local uploads 또는 S3에 파일을 저장하고 publicUrl/objectKey를 받습니다.",
        "ImageIO 기반 사전 검사로 부적합 이미지는 NOT_APPLICABLE 처리합니다.",
        "IDeepfakeClient 구현체가 Reality Defender, IMD, Dummy 중 설정에 맞게 분석합니다.",
        "Reality Defender real 모드에서는 IMD 히트맵을 추가 호출해 processedHeatmap, overlay, regions를 병합합니다.",
        "IVerifyMapper.insertVerification()으로 verification 테이블에 저장합니다.",
        "ResultController가 /detail/{id} 또는 /result/{id}에서 권한을 확인하고 deepfake-result.jsp로 결과를 보여줍니다.",
    ]:
        numbered(doc, step)

    h(doc, "회원/로그인 기능 흐름", 1)
    for step in [
        "/login은 login.jsp, /signup은 signup.jsp, /find-account는 find-account.jsp, /mypage는 mypage.jsp를 반환합니다.",
        "AuthApiController.login()은 회원을 조회하고 BCrypt로 비밀번호를 검증합니다.",
        "기존 SHA-256 비밀번호 해시가 남아 있으면 로그인 성공 시 BCrypt로 재저장합니다.",
        "로그인과 Google OAuth 성공 시 request.changeSessionId()로 세션 ID를 교체합니다.",
        "회원가입은 이메일 중복 확인, 인증번호 발송, 인증번호 검증, 최종 저장 순서입니다.",
        "Google OAuth는 GoogleOAuthController가 직접 인증 URL 생성, state 검증, token/userinfo 요청, 회원 연결을 처리합니다.",
    ]:
        numbered(doc, step)

    h(doc, "뉴스/커뮤니티/기록 기능 흐름", 1)
    for step in [
        "NewsController는 /news와 /news/{id}를 처리하고, NewsService는 Naver News API, MongoDB 캐시, 로컬 메모리 캐시를 조합합니다.",
        "CommunityController는 community.jsp/create-post.jsp/post-detail.jsp 화면을 반환하고, CommunityApiController가 게시글·댓글·좋아요 REST API를 처리합니다.",
        "HistoryController는 /history에서 deepfake-history.jsp를 반환하고 로그인 사용자의 검증 기록을 조회합니다.",
        "VerificationApiController는 검증 생성, 목록 조회, 상세 조회, 삭제 API를 제공하고 userId 조건으로 본인 기록만 삭제합니다.",
    ]:
        numbered(doc, step)

    h(doc, "주요 설정 파일", 1)
    file_section(doc, "Config/초기화", [
        ("src/main/java/kopo/poly/Deepfake2Application.java", "Spring Boot 시작 클래스입니다. @MapperScan으로 MyBatis Mapper를 등록합니다."),
        ("src/main/java/kopo/poly/config/AppConfig.java", "RestClient 공통 Bean을 만들고 HTTP connect/read timeout을 설정합니다."),
        ("src/main/java/kopo/poly/config/WebConfig.java", "CORS와 /uploads/** 정적 리소스 매핑을 설정합니다."),
        ("src/main/java/kopo/poly/config/OriginProtectionFilter.java", "POST/PUT/PATCH/DELETE 요청의 Origin/Referer를 검사해 외부 origin의 상태 변경 요청을 차단합니다."),
        ("src/main/java/kopo/poly/config/AppUserSchemaInitializer.java", "app_user와 커뮤니티 관련 테이블/컬럼을 보정합니다."),
        ("src/main/java/kopo/poly/config/EmailAuthSchemaInitializer.java", "email_auth 테이블과 인덱스를 준비합니다."),
        ("src/main/java/kopo/poly/config/VerificationSchemaInitializer.java", "verification 테이블의 user_id, api_raw, 인덱스를 보정합니다."),
        ("src/main/java/kopo/poly/config/DetectionSchemaInitializer.java", "verification_record 테이블을 준비합니다."),
        ("src/main/resources/application.properties", "JSP 경로, 포트, multipart, 저장소, DB, Mail, Reality Defender, IMD, Naver, MongoDB, Google OAuth 설정을 관리합니다."),
        ("build.gradle", "Gradle 빌드, Java 17, bootWar, MyBatis, Web, Mail, MongoDB, Security Crypto, AWS S3, JSP/JSTL 의존성을 정의합니다."),
    ])

    file_section(doc, "화면 Controller", [
        ("HomeController.java", "/와 /home 요청을 home.jsp로 연결합니다."),
        ("AuthPageController.java", "/signup, /login, /find-account, /mypage 화면을 반환합니다."),
        ("UploadController.java", "JSP 업로드 요청 /upload를 받아 VerifyService로 분석하고 /detail/{id}로 redirect합니다."),
        ("ResultController.java", "/detail/{id}, /result/{id} 결과 화면을 처리하고 deepfake-result.jsp로 모델을 전달합니다."),
        ("HistoryController.java", "/history에서 사용자 검증 기록을 조회하고 deepfake-history.jsp를 반환합니다."),
        ("NewsController.java", "/news 목록, /news/refresh, /news/{id} 상세 화면을 처리합니다."),
        ("CommunityController.java", "/community, /community/create, /community/{id} 화면을 반환합니다."),
        ("GoogleOAuthController.java", "Google OAuth 인증 시작과 콜백을 수동으로 처리합니다."),
        ("FaqController.java", "/faq 화면을 반환합니다."),
        ("ReportController.java", "/report 및 /report/success 안내 화면을 처리합니다."),
        ("RealityDefenderController.java", "별도 /detect 흐름의 Reality Defender 분석 화면 진입을 처리합니다."),
    ])

    file_section(doc, "API Controller", [
        ("controller/api/AuthApiController.java", "로그인, 회원가입, 이메일 인증 연계, 비밀번호 재설정, 프로필 수정, 비밀번호 변경, 로그아웃, 회원 탈퇴 API를 처리합니다."),
        ("controller/api/VerificationApiController.java", "이미지 검증 생성, 검증 목록/상세 조회, 검증 기록 삭제 API를 제공합니다."),
        ("controller/api/CommunityApiController.java", "게시글, 댓글, 대댓글, 좋아요 REST API를 제공합니다."),
        ("controller/api/EmailAuthController.java", "이메일 인증번호 발송과 검증 API를 제공합니다."),
    ])

    file_section(doc, "Service와 Client", [
        ("VerifyService.java", "메인 분석 파이프라인입니다. 파일 검증, 저장소 업로드, 사전 이미지 검사, 외부 분석 호출, 히트맵 병합, DB 저장, 상세 분석 JSON 생성을 담당합니다."),
        ("RealityDefenderClient.java", "Reality Defender presigned upload, 파일 업로드, 결과 polling을 수행합니다."),
        ("ImdDeepfakeClient.java", "로컬 IMD FastAPI /analyze에 파일을 보내 분석 결과를 받습니다."),
        ("ImdHeatmapClient.java", "IMD 분석 결과 중 heatmap/overlay/regions를 추출해 상세 화면 시각화를 보강합니다."),
        ("ImdLocalServiceManager.java", "로컬 IMD 서버 health check와 선택적 자동 시작을 담당합니다."),
        ("DummyDeepfakeClient.java", "API 없이 테스트용 분석 결과를 반환합니다."),
        ("DummyObjectStorageService.java", "local 저장 모드에서 uploads 폴더에 파일을 저장합니다."),
        ("S3ObjectStorageService.java", "s3 저장 모드에서 Amazon S3에 파일을 업로드합니다."),
        ("NewsService.java", "Naver News API 호출, 기사 정제, 딥페이크 관련 필터링, 중복 제거, MongoDB/메모리 캐시를 담당합니다."),
        ("CommunityService.java", "게시글/댓글 작성자 권한 확인, 조회수 증가, 좋아요 토글, 댓글 트리 구성을 담당합니다."),
        ("EmailAuthService.java", "이메일 인증번호 생성, 저장, 발송, 검증, 만료 여부 확인을 담당합니다."),
        ("MailService.java", "JavaMailSender로 실제 메일을 발송합니다."),
        ("HistoryService.java", "로그인 사용자의 검증 기록 목록을 조회합니다."),
        ("RealityDefenderService.java", "별도 Reality Defender detect 흐름의 분석과 기록 저장을 처리합니다."),
    ])

    file_section(doc, "Mapper/XML/DB", [
        ("IUserMapper.java + UserMapper.xml", "app_user 회원 조회/가입/소셜 연결/프로필 수정/비밀번호 변경/회원 탈퇴 관련 SQL입니다."),
        ("IVerifyMapper.java + VerifyMapper.xml", "verification 검증 결과 insert/select/list/delete SQL입니다."),
        ("ICommunityMapper.java + CommunityMapper.xml", "community_post, community_comment, like 테이블 CRUD와 조회수/좋아요 SQL입니다."),
        ("IEmailAuthMapper.java + EmailAuthMapper.xml", "email_auth 인증번호 저장, 최신 인증 조회, 인증 완료 처리 SQL입니다."),
        ("IHistoryMapper.java + HistoryMapper.xml", "사용자별 검증 기록 목록 조회 SQL입니다."),
        ("IRealityDefenderMapper.java + RealityDefenderMapper.xml", "verification_record 별도 탐지 기록 저장/조회 SQL입니다."),
        ("NewsCacheDocument.java + NewsCacheRepository.java", "MongoDB news_cache 컬렉션의 뉴스 캐시 문서와 Repository입니다."),
    ])

    file_section(doc, "DTO/Util", [
        ("ApiResponse.java", "REST API 응답을 success/data/error 구조로 통일합니다."),
        ("VerifyDTO.java, VerificationCreateResponseDTO.java, DeepfakeResultDTO.java", "검증 생성·저장·분석 결과 전달 DTO입니다."),
        ("LoginRequestDTO.java, SignupCompleteRequestDTO.java, ResetPasswordRequestDTO.java", "회원 API 요청값을 받는 DTO입니다."),
        ("CommunityPostDTO.java, CommunityCommentDTO.java", "커뮤니티 게시글/댓글 화면과 API 응답 DTO입니다."),
        ("NewsViewDTO.java, NaverNewsResponseDTO.java, NaverNewsItemDTO.java", "뉴스 API 응답과 상세 화면 전달 DTO입니다."),
        ("SessionUtil.java", "세션 USER_ID 추출과 검증 결과 접근 권한 확인을 담당합니다."),
        ("HashUtil.java", "기존 SHA-256 비밀번호 해시 검증을 위한 legacy 유틸입니다."),
        ("CommonUtil.java", "문자열 null 처리 등 공통 보조 기능을 제공합니다."),
    ])

    file_section(doc, "JSP 화면", [
        ("home.jsp", "메인 업로드 화면입니다. 파일 선택, 미리보기, 분석 요청을 처리합니다."),
        ("login.jsp / signup.jsp / find-account.jsp / mypage.jsp", "로그인, 회원가입, 계정 찾기, 마이페이지 화면입니다."),
        ("deepfake-history.jsp / deepfake-result.jsp / analyzing.jsp", "검증 기록, 검증 결과, 분석 중 화면입니다."),
        ("news.jsp / news-detail.jsp", "뉴스 목록과 뉴스 상세 화면입니다."),
        ("community.jsp / create-post.jsp / post-detail.jsp", "커뮤니티 목록, 작성, 상세 화면입니다."),
        ("faq.jsp / report.jsp", "FAQ와 신고 안내 화면입니다."),
        ("deepfake-upload.jsp 등 새 이름 JSP", "최근 UI 전환 과정에서 남아 있는 새 이름의 JSP입니다. 현재 일부 컨트롤러는 기존 이름과 새 이름을 함께 사용하므로 실제 반환 view name을 Controller 기준으로 확인해야 합니다."),
    ])

    file_section(doc, "IMD/배포", [
        ("imd-service/app.py", "Spring Boot가 호출하는 로컬 FastAPI 기반 IMD 분석/히트맵 서비스입니다."),
        ("imd-service/start.ps1", "Windows 로컬에서 IMD 서비스를 실행하는 스크립트입니다."),
        ("deploy/docker-compose.imd.prod.yml", "운영 환경에서 IMD 서비스를 컨테이너로 실행하기 위한 compose 파일입니다."),
        ("deploy/aws-deploy-guide.md", "AWS 배포 절차와 환경변수 설정 가이드입니다."),
        ("deploy/aws-env.properties.example", "운영 .env.properties 예시입니다."),
    ])

    h(doc, "운영 시 주의할 설정", 1)
    for line in [
        "APP_STORAGE_TYPE=local이면 uploads 폴더를 사용하고, APP_STORAGE_TYPE=s3이면 AWS S3 자격증명 또는 EC2 IAM Role이 필요합니다.",
        "DEEPFAKE_CLIENT_MODE=real이면 Reality Defender가 주 분석이고 IMD는 히트맵 보강에 사용됩니다.",
        "IMD_HEATMAP_ENABLED=true이면 분석 결과에 시각화 히트맵을 추가하려고 시도합니다.",
        "APP_ALLOWED_ORIGINS는 OriginProtectionFilter에서 허용할 외부 origin 목록입니다.",
        "GOOGLE_CLIENT_ID/GOOGLE_CLIENT_SECRET이 disabled 또는 비어 있으면 Google OAuth 진입 시 로그인 오류로 돌아갑니다.",
    ]:
        bullet(doc, line)

    doc.save(FULL_OUT)


def main():
    DOCS.mkdir(exist_ok=True)
    build_qna()
    build_full()


if __name__ == "__main__":
    main()
