from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor


SOURCE = r"C:\Users\8316-22\Downloads\딥페이크_프로젝트_예상질문_답변_정리본_자연스러운버전.docx"
OUTPUT = r"C:\SpringBootWorks\deepfake2\docs\딥페이크_프로젝트_예상질문_답변_보강_코드설명.docx"


supplements = [
    (
        "전체 구조 설명 보강",
        "발표에서는 kopo.poly 아래를 기준으로 controller, controller.api, service, service.impl, mapper, dto, config, document, repository로 나눠 설명하면 좋습니다. "
        "화면 이동은 JSP용 Controller가 담당하고, 비동기 요청이나 JSON 응답은 controller.api가 담당합니다. 실제 처리 흐름은 Service에서 묶고, MariaDB는 MyBatis Mapper, MongoDB 뉴스 캐시는 NewsCacheRepository로 접근합니다."
    ),
    (
        "핵심 기능 흐름 보강",
        "이미지 판별 기능은 UploadController 또는 VerificationApiController에서 파일을 받고 VerifyService.createVerification()으로 넘기는 구조입니다. "
        "VerifyService는 이미지 형식과 용량을 검사하고, 저장소에 파일을 저장한 뒤 Reality Defender 또는 IMD 분석 클라이언트를 호출합니다. "
        "분석 결과는 verdict, score, apiProvider, apiRaw 형태로 verification 테이블에 저장하고, 상세 화면에서는 점수, 판정, 히트맵, 설명 문구로 풀어서 보여줍니다."
    ),
    (
        "인증/회원 보강",
        "인증 기능은 AuthApiController가 로그인, 회원가입, 아이디 찾기, 비밀번호 재설정, 프로필 수정, 회원 탈퇴 API를 담당합니다. "
        "회원가입은 이메일 중복 확인, 인증번호 발송, 인증번호 검증, 최종 가입 순서로 진행됩니다. 비밀번호는 HashUtil.sha256()으로 해시 처리해서 저장하고, 로그인 성공 시 HttpSession에 USER_ID, USER_NAME, USER_EMAIL을 저장합니다."
    ),
    (
        "외부 API 보강",
        "외부 API는 Controller에 직접 넣지 않고 Client 또는 Service로 분리했습니다. RealityDefenderClient는 실제 딥페이크 분석 API를 호출하고, ImdDeepfakeClient와 ImdHeatmapClient는 로컬 Python IMD 서비스와 통신합니다. "
        "NewsService는 Naver News API를 호출한 뒤 MongoDB와 메모리 캐시에 저장해서 반복 호출을 줄입니다. Google 로그인은 GoogleOAuthController와 GoogleOAuthService에서 OAuth 시작, 콜백, 회원 연결, 세션 저장을 처리합니다."
    ),
    (
        "DB/저장소 보강",
        "MariaDB는 app_user, verification, email_auth, community_post, community_comment 같은 업무 데이터를 저장합니다. "
        "MongoDB는 뉴스 캐시처럼 문서 형태로 저장하기 좋은 데이터를 담당합니다. 파일 저장은 로컬 uploads 또는 S3 저장소로 바꿀 수 있게 IObjectStorageService 인터페이스와 구현체를 나눴습니다."
    ),
    (
        "예외 처리/보안 보강",
        "API 응답은 ApiResponse.ok(), ApiResponse.fail() 형태로 통일했습니다. 입력값이 잘못되면 400 계열 코드, 로그인이 필요하면 401, 권한이 없으면 403, 데이터가 없으면 404, 서버 오류는 500 계열 코드로 구분합니다. "
        "세션의 userId와 데이터의 userId를 비교해서 다른 사용자의 검증 기록이나 게시글을 임의로 조회, 수정, 삭제하지 못하게 막았습니다."
    ),
]


sections = [
    ("설계", [
        ("메뉴 구조도", "JSP 화면과 URL을 기준으로 사용자가 어떤 메뉴에서 어떤 기능으로 이동하는지 정리한 부분입니다. 홈, 업로드/분석, 결과, 기록, 뉴스, 커뮤니티, FAQ, 신고, 마이페이지, 로그인/회원가입 메뉴 흐름을 설명하면 됩니다."),
        ("프로그램 명세서", "프로젝트가 제공하는 기능을 사용자 관점에서 정리한 문서입니다. 예를 들면 이미지 업로드, 딥페이크 판별, 결과 상세 조회, 뉴스 조회, 게시판, 회원 기능이 어떤 입력과 출력을 갖는지 설명합니다."),
        ("컬렉션 정의서(NoSQL)", "MongoDB에 저장되는 news_cache 컬렉션 구조를 설명합니다. NewsCacheDocument가 제목, 설명, 링크, 언론사, 발행일, 캐시 생성일, 만료일, 검색어 정보를 저장합니다."),
        ("테이블 명세서(RDBMS)", "MariaDB 테이블 구조를 설명합니다. app_user, email_auth, verification, verification_record, community_post, community_comment 같은 테이블의 컬럼과 역할을 정리합니다."),
        ("논리 ERD 작성", "업무 관계를 중심으로 회원, 검증 기록, 게시글, 댓글, 이메일 인증 정보가 어떻게 연결되는지 설명합니다. 예를 들어 회원 1명이 여러 검증 기록과 게시글을 가질 수 있습니다."),
        ("물리 ERD 작성", "실제 DB 컬럼명, 자료형, PK/FK, 인덱스 기준으로 MariaDB 구조를 설명합니다. MyBatis Mapper XML에서 사용하는 SQL과 연결해서 말하면 좋습니다."),
        ("API 연동 설계", "프론트 요청 URL, HTTP Method, 요청 DTO, 응답 DTO, 호출 Service를 정리합니다. 예시는 /api/v1/verifications, /api/v1/auth/login, /api/v1/community/posts, /news입니다."),
        ("화면 흐름 검토", "사용자가 로그인 후 업로드, 분석 대기, 결과 확인, 기록 조회로 이동하는 흐름을 검토합니다. 실패 시에는 에러 메시지 또는 재시도 화면으로 돌아가는 흐름도 함께 설명합니다."),
    ]),
    ("개발환경 세팅", [
        ("Spring Boot 세팅", "Deepfake2Application에서 Spring Boot가 시작되고 @MapperScan으로 MyBatis Mapper를 스캔합니다. application.properties에서 JSP 경로, 포트, multipart, DB, 메일, 외부 API 설정을 관리합니다."),
        ("MariaDB 세팅", "spring.datasource.* 설정으로 MariaDB에 연결합니다. app_user, verification, community 관련 테이블은 MyBatis Mapper XML을 통해 조회, 저장, 수정, 삭제합니다."),
        ("MongoDB 세팅", "spring.data.mongodb.uri로 MongoDB에 연결합니다. NewsCacheRepository가 MongoRepository를 상속해서 news_cache 컬렉션의 뉴스 캐시를 조회하고 저장합니다."),
    ]),
    ("구현(인증/회원)", [
        ("로그인 화면 진입", "AuthPageController 또는 로그인 JSP가 로그인 화면을 반환합니다. 사용자가 이메일과 비밀번호를 입력하면 /api/v1/auth/login으로 요청합니다."),
        ("아이디/비밀번호 입력", "LoginRequestDTO로 email과 password를 받습니다. Controller에서 null, 공백 여부를 검사한 뒤 userMapper.selectUserByEmail()로 회원을 조회합니다."),
        ("입력값 유효성 검사", "이메일, 비밀번호, 이름, 전화번호, 주소처럼 필수값은 AuthApiController에서 서버 측 검증을 다시 수행합니다. 프론트 검증을 우회해도 서버에서 막기 위한 처리입니다."),
        ("로그인 요청/성공 처리", "입력 비밀번호를 HashUtil.sha256()으로 해시한 값과 DB의 password_hash를 비교합니다. 성공하면 세션에 USER_ID, USER_NAME, USER_EMAIL을 저장합니다."),
        ("구글 로그인 연동", "GoogleOAuthController가 구글 인증 시작과 콜백을 처리하고, GoogleOAuthService가 구글 사용자 정보를 조회해 세션 저장 또는 회원 정보 연결을 수행합니다."),
        ("회원가입 화면 진입", "회원가입 JSP에서 이메일, 인증번호, 비밀번호, 이름, 전화번호, 주소를 입력합니다. 실제 가입은 중복 확인, 인증번호 발송, 인증 확인, 최종 가입 순서로 분리됩니다."),
        ("아이디 중복 체크", "/api/v1/auth/signup/check-email에서 userMapper.existsByEmail()을 호출해 이미 가입된 이메일인지 확인합니다."),
        ("비밀번호 일치 확인", "프론트에서 1차로 새 비밀번호와 확인값을 비교하고, 서버에서는 최종 가입 또는 재설정 시 비밀번호 길이 같은 필수 조건을 다시 확인합니다."),
        ("이메일 인증번호 발송", "EmailAuthService.sendAuthCode()가 인증번호를 만들고 MailService를 통해 메일을 발송합니다. 인증 정보는 email_auth 테이블에 저장됩니다."),
        ("인증번호 확인", "EmailAuthService.verifyAuthCode()가 이메일과 코드가 맞는지, 만료되지 않았는지 확인하고 verified 상태로 변경합니다."),
        ("회원가입 처리", "/api/v1/auth/signup/complete에서 인증 완료 여부를 확인한 뒤 비밀번호를 SHA-256으로 해시해서 app_user 테이블에 INSERT합니다."),
        ("아이디 찾기 화면 진입", "아이디 찾기 화면에서 이름과 전화번호를 입력받고 /api/v1/auth/find-id로 요청합니다."),
        ("아이디 찾기 이메일 인증번호 발송", "현재 프로젝트에서는 아이디 찾기는 이름과 전화번호 기준으로 조회하고, 이메일 인증은 회원가입/비밀번호 재설정 흐름에서 사용합니다. 발표 때는 이 차이를 분명히 말하면 좋습니다."),
        ("아이디 찾기 인증번호 확인", "아이디 찾기 자체에는 인증번호 검증보다 회원 정보 조회가 핵심입니다. 이메일 인증 확인 로직은 EmailAuthService의 verifyAuthCode()를 재사용할 수 있는 구조입니다."),
        ("아이디 조회 결과 출력", "userMapper.selectUserByNameAndPhone() 결과로 이메일과 이름을 반환합니다. 사용자는 화면에서 가입 이메일을 확인합니다."),
        ("비밀번호 재설정 화면 진입", "비밀번호 찾기 화면에서 이메일과 이름을 입력하고 인증번호 발송을 요청합니다."),
        ("계정/이메일 일치 확인", "/api/v1/auth/password/send-code에서 이메일과 이름으로 userMapper.selectUserByEmailAndName()을 조회해 실제 회원인지 확인합니다."),
        ("인증번호 확인", "/api/v1/auth/password/verify-code에서 EmailAuthService.verifyAuthCode()로 인증번호를 검증합니다."),
        ("새 비밀번호 입력/검증", "/api/v1/auth/password/reset에서 새 비밀번호 길이를 확인하고 이메일 인증 완료 여부를 확인합니다."),
        ("비밀번호 변경 처리", "HashUtil.sha256()으로 새 비밀번호를 해시한 뒤 userMapper.updatePasswordByEmail()로 DB를 수정합니다."),
    ]),
    ("구현(딥페이크 판별)", [
        ("이미지 업로드", "UploadController 또는 VerificationApiController가 multipart/form-data로 넘어온 file을 받습니다. 받은 파일은 VerifyService.createVerification()으로 전달됩니다."),
        ("이미지 파일 선택", "deepfake-upload.jsp에서 파일 input을 통해 이미지를 선택합니다. 서버는 request.getFile(\"file\") 이름으로 파일을 받습니다."),
        ("파일 형식/용량 검사", "VerifyService.validateImageOnly()에서 MIME 타입이 image/로 시작하는지, 확장자가 png/jpg/jpeg/bmp/webp/heic/heif인지, 크기가 32MB 이하인지 확인합니다."),
        ("이미지 미리보기 출력", "업로드 화면의 JS가 사용자가 선택한 이미지를 FileReader 등으로 미리 보여줍니다. 서버에서는 저장 후 publicUrl을 내려줘 결과 화면에서도 원본 이미지를 표시합니다."),
        ("업로드 데이터 생성", "날짜와 UUID를 조합해 objectKey를 만들고 IObjectStorageService.uploadPublic()으로 로컬 uploads 또는 S3에 저장합니다."),
        ("분석 요청 API 호출", "VerifyService가 IDeepfakeClient.analyze()를 호출합니다. 구현체는 설정에 따라 RealityDefenderClient, ImdDeepfakeClient, DummyDeepfakeClient로 바뀔 수 있습니다."),
        ("로딩/진행 상태 표시", "분석 요청 후 deepfake-analyzing.jsp 또는 프론트 JS에서 결과가 나올 때까지 로딩 상태를 보여줍니다. 분석이 끝나면 결과 상세 화면으로 이동합니다."),
        ("응답 수신 처리", "DeepfakeResultDTO로 verdict, score, raw JSON을 받습니다. Reality Defender 결과와 IMD 히트맵 결과가 있으면 raw JSON에 합쳐 저장합니다."),
        ("분석 실패 예외 처리", "이미지 조건이 맞지 않으면 NOT_APPLICABLE 결과를 만들고, 외부 API 오류는 try-catch로 잡아 ApiResponse.fail() 또는 화면 에러 메시지로 처리합니다."),
        ("결과 화면 표시", "ResultController/DetailController가 검증 ID를 기준으로 VerifyService.getOne()을 호출하고, DB에 저장된 판정 결과와 이미지 URL을 JSP에 전달합니다."),
        ("판별 점수 출력", "VerifyDTO의 score를 화면에서 퍼센트 또는 위험도 형태로 보여줍니다. score가 없거나 분석 불가인 경우에는 별도 안내 문구를 표시합니다."),
        ("진위 여부/요약 문구 출력", "verdict 값과 enhancedExplanation을 이용해 정상, 주의, 의심, 강한 의심, 분석 불가 같은 사용자 친화적인 문구를 출력합니다."),
        ("분석 이미지 표시(IMD 오픈소스)", "IMD Heatmap API가 제공한 processedHeatmap, overlayHeatmap, suspicious regions 정보를 상세 화면에서 시각화합니다."),
        ("상세 보기 버튼 처리", "목록이나 결과 화면의 상세 보기 버튼은 검증 ID를 URL 또는 API 요청으로 넘겨 DetailController/VerificationApiController.detail()을 호출합니다."),
        ("세부 분석 정보 출력", "api_raw를 파싱해 분석 출처, 위치 정보, 히트맵, 증상 설명, 주의사항, 추천 조치 등을 화면에 보여줍니다."),
        ("상세 점수/항목 표시", "상세 화면에서는 점수뿐 아니라 판정 근거, 의심 영역 개수, 히트맵 존재 여부, 분석 제공자까지 함께 표시할 수 있습니다."),
    ]),
    ("구현(신고)", [
        ("신고 페이지 진입", "ReportController가 /report 요청을 받아 report.jsp를 반환합니다."),
        ("맵 출력 API 호출", "신고 안내 화면에서 Kakao Map 같은 지도 API를 사용한다면 프론트에서 스크립트를 로딩해 경찰서 또는 신고 기관 위치를 표시합니다."),
        ("경찰서 마커 표시", "사용자 위치나 지정 좌표 주변에 경찰서 마커를 표시하고, 신고가 필요한 경우 실제 기관으로 이동할 수 있게 안내합니다."),
    ]),
    ("구현(자주하는질문)", [
        ("FAQ 목록 조회", "FaqController가 FAQ 화면을 반환하고, faq.jsp에서 질문 목록을 출력합니다."),
        ("질문/답변 펼치기", "FAQ 화면의 JS 또는 CSS로 아코디언 UI를 구성해 질문을 클릭하면 답변이 펼쳐지게 처리합니다."),
    ]),
    ("구현(뉴스)", [
        ("뉴스 메인 화면 진입", "NewsController의 /news가 news-list.jsp를 반환합니다. 기본 검색어는 딥페이크이고, 카테고리, 정렬, 페이지 값을 함께 처리합니다."),
        ("카테고리별 뉴스 조회", "제목과 설명의 키워드를 기준으로 범죄사례, 법률/규제, 기술동향, 피해사례, 대응방법 같은 카테고리를 추정해 필터링합니다."),
        ("뉴스 검색 API 호출", "NewsService.requestNaverNews()가 Naver News API에 GET 요청을 보내고, X-Naver-Client-Id와 X-Naver-Client-Secret을 헤더에 담습니다."),
        ("뉴스 목록 출력", "NewsService가 API 결과를 정제, 중복 제거, MongoDB 캐시 저장 후 NewsController가 JSP 모델에 newsList를 담아 목록을 출력합니다."),
        ("뉴스 자세히 보기", "/news/{id}에서 MongoDB 또는 로컬 캐시의 기사 정보를 찾아 news-detail.jsp로 넘깁니다."),
        ("관련 링크 이동", "news-detail.jsp에서 originallink 또는 link를 사용해 원문 기사로 이동할 수 있게 처리합니다."),
    ]),
    ("구현(커뮤니티)", [
        ("게시글 목록 조회", "CommunityApiController의 GET /api/v1/community/posts가 목록을 반환합니다. keyword가 있으면 검색 조건으로 넘깁니다."),
        ("게시글 검색/정렬", "ICommunityMapper.selectPosts()와 CommunityMapper.xml SQL에서 검색어 기준 조회를 수행합니다. 화면에서는 최신순, 조회수, 좋아요 같은 기준으로 확장 가능합니다."),
        ("게시글 작성", "POST /api/v1/community/posts에서 세션 USER_ID와 USER_NAME을 확인하고 제목, 내용, topic을 검증한 뒤 CommunityService.createPost()를 호출합니다."),
        ("게시글 수정", "PUT /api/v1/community/posts/{id}에서 현재 로그인 사용자와 게시글 작성자가 같은지 확인한 뒤 updatePost()를 실행합니다."),
        ("게시글 상세 조회", "GET /api/v1/community/posts/{id}가 상세 데이터를 반환하고, CommunityService.getPost()는 조회수 증가 후 게시글을 조회합니다."),
        ("댓글 목록 조회", "GET /api/v1/community/posts/{id}/comments가 댓글 목록을 반환합니다. Service에서 parentId 기준으로 댓글과 대댓글 트리 구조를 만듭니다."),
        ("댓글 작성", "POST /comments에서 로그인 여부와 내용 공백 여부를 확인한 뒤 insertComment를 실행합니다. parentId가 있으면 대댓글로 처리합니다."),
        ("댓글 수정/삭제", "PUT/DELETE /comments/{commentId}에서 댓글 작성자와 현재 사용자 ID가 같은지 확인한 뒤 수정 또는 삭제합니다."),
    ]),
    ("구현(검증기록)", [
        ("검증기록 목록 조회", "VerificationApiController의 GET /api/v1/verifications 또는 HistoryController가 로그인 사용자의 검증 기록을 조회합니다. userId를 조건으로 걸어 본인 기록만 보여줍니다."),
        ("검증기록 상세 조회", "GET /api/v1/verifications/{id} 또는 상세 화면에서 VerifyService.getOne()으로 DB 데이터를 가져옵니다. SessionUtil.canAccessVerification()으로 접근 권한을 확인합니다."),
    ]),
    ("구현(마이페이지)", [
        ("회원정보 수정", "PUT /api/v1/auth/profile에서 세션 userId를 확인하고 이름, 전화번호, 주소를 검증한 뒤 userMapper.updateProfileById()를 실행합니다."),
        ("회원 탈퇴", "DELETE /api/v1/auth/account에서 사용자의 좋아요, 댓글, 게시글, 검증 기록을 순서대로 삭제한 뒤 app_user를 삭제하고 세션을 무효화합니다."),
        ("로그아웃", "POST /api/v1/auth/logout에서 session.invalidate()를 호출해 서버에 저장된 로그인 상태를 제거합니다."),
        ("프로그램 점검", "발표 때는 DB 연결, 외부 API 키, 메일 계정, MongoDB URI, IMD 로컬 서비스 실행 여부, 파일 업로드 경로 권한을 점검 항목으로 말하면 좋습니다."),
    ]),
    ("배포/유지보수", [
        ("배포 스크립트", "deploy 폴더와 docker-compose.imd.yml을 이용해 애플리케이션, DB, IMD 서비스 실행 환경을 맞출 수 있습니다. 운영에서는 환경변수로 DB/API Key를 주입하는 방식이 안전합니다."),
        ("유지보수", "Controller는 요청 처리, Service는 기능 흐름, Mapper/Repository는 저장소 접근으로 역할이 나뉘어 있어 기능을 수정할 때 영향 범위를 찾기 쉽습니다. API Key, DB 설정, 파일 저장 방식은 application.properties와 환경변수로 관리합니다."),
    ]),
]


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "맑은 고딕"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    set_run_font(r, 16 if level == 1 else 13, True, (31, 78, 121) if level == 1 else (68, 68, 68))
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_bullet(doc, title, body):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"{title}: ")
    set_run_font(r1, 10.2, True)
    r2 = p.add_run(body)
    set_run_font(r2, 10.2)
    return p


def main():
    doc = Document(SOURCE)

    for section in doc.sections:
        section.top_margin = section.top_margin

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    add_heading(doc, "기존 답변에 추가하면 좋은 보강 답변", 1)
    add_body(doc, "아래 내용은 기존 예상 질문 답변 뒤에 덧붙여 말하면 좋은 보충 설명입니다. 실제 프로젝트 클래스명과 처리 흐름을 중심으로 정리했습니다.")
    for title, body in supplements:
        add_bullet(doc, title, body)

    add_heading(doc, "코드별 발표 설명 정리", 1)
    add_body(doc, "왼쪽 체크리스트 항목 기준으로, 발표 중 코드를 보여주면서 설명하기 좋은 내용을 정리했습니다. 기본 설명 순서는 화면/JSP, Controller/API, Service, Mapper 또는 Repository, DB/외부 API입니다.")

    for section_title, items in sections:
        add_heading(doc, section_title, 2)
        for title, body in items:
            add_bullet(doc, title, body)

    add_heading(doc, "발표 때 말하기 좋은 한 줄 흐름", 1)
    flow_lines = [
        "로그인: login.jsp에서 입력한 값을 AuthApiController가 받고, IUserMapper로 회원을 조회한 뒤 세션에 로그인 정보를 저장합니다.",
        "회원가입: 이메일 중복 확인과 인증번호 검증을 통과한 사용자만 app_user 테이블에 저장합니다.",
        "딥페이크 분석: 업로드 파일을 VerifyService가 검증, 저장, 외부 분석 API 호출, 결과 DB 저장까지 한 번에 처리합니다.",
        "결과 상세: 저장된 api_raw와 score를 다시 해석해서 판정 근거, 히트맵, 주의사항을 화면에 보여줍니다.",
        "뉴스: Naver News API 결과를 MongoDB에 캐시해서 반복 요청을 줄이고, 화면에서는 카테고리와 검색 조건으로 보여줍니다.",
        "커뮤니티: 게시글과 댓글은 세션 userId로 작성자 권한을 확인한 뒤 수정/삭제를 허용합니다.",
    ]
    for line in flow_lines:
        add_body(doc, line)

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
