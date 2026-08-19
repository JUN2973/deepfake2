import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "docs" / "로그인_회원가입_JSP부터_Java까지_설명.md"
OUT = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else ROOT / "docs" / "로그인_회원가입_JSP부터_Java까지_설명.docx"


def set_run_font(run, name="맑은 고딕", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_code_line(doc, line):
    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=0, line=1.0)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.05)
    shade_paragraph(p, "F3F6FA")
    run = p.add_run(line if line else " ")
    set_run_font(run, name="Consolas", size=8.5, color="1F2937")


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color="1F2937")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color="1F2937")


def add_heading(doc, text, level):
    style_name = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
    p = doc.add_paragraph(style=style_name)
    if level == 1:
        style_paragraph(p, before=18, after=10)
        size = 16
        color = "2E74B5"
    elif level == 2:
        style_paragraph(p, before=14, after=7)
        size = 13
        color = "2E74B5"
    else:
        style_paragraph(p, before=10, after=5)
        size = 12
        color = "1F4D78"
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)


def normalize_inline(text):
    return text.replace("`", "")


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("로그인 / 회원가입 JSP부터 Java까지 연결 설명")
    set_run_font(title_run, size=20, bold=True, color="1F4D78")

    meta = doc.add_paragraph()
    style_paragraph(meta, after=12)
    meta_run = meta.add_run("대상 파일: login.jsp, signup.jsp, AuthPageController, AuthApiController, AuthService, EmailAuthService, Mapper XML")
    set_run_font(meta_run, size=9.5, color="64748B")

    in_code = False
    code_gap_pending = False
    paragraph_buffer = []

    def flush_paragraph():
        if paragraph_buffer:
            text = " ".join(paragraph_buffer).strip()
            if text:
                add_body_paragraph(doc, normalize_inline(text))
            paragraph_buffer.clear()

    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            in_code = not in_code
            if not in_code:
                code_gap_pending = True
            continue

        if in_code:
            add_code_line(doc, line)
            continue

        if code_gap_pending and stripped:
            spacer = doc.add_paragraph()
            style_paragraph(spacer, after=2)
            code_gap_pending = False

        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            add_heading(doc, stripped[2:].strip(), 1)
        elif stripped.startswith("## "):
            flush_paragraph()
            add_heading(doc, stripped[3:].strip(), 2)
        elif stripped.startswith("### "):
            flush_paragraph()
            add_heading(doc, stripped[4:].strip(), 3)
        elif stripped.startswith("- "):
            flush_paragraph()
            add_bullet(doc, normalize_inline(stripped[2:].strip()))
        else:
            paragraph_buffer.append(stripped)

    flush_paragraph()

    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer_run = footer.add_run("DeepScan 로그인/회원가입 코드 설명")
    set_run_font(footer_run, size=8.5, color="64748B")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_docx()
